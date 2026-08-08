from openpyxl import Workbook

wb = Workbook()
ws = wb.active
ws.title = "TESTE"

ws["A1"] = "CODIGO"
ws["B1"] = "PRODUTO"
ws["C1"] = "ESTOQUE"

ws["A2"] = 1001
ws["B2"] = "Dipirona"
ws["C2"] = 25

ws["A3"] = 1002
ws["B3"] = "Paracetamol"
ws["C3"] = 40

wb.save("teste_excel.xlsx")

print("Planilha criada com sucesso!")