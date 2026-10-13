# -*- coding: utf-8 -*-
"""
EXTRATOR BASE FISCAL XML - PROJETO XML / BASE FISCAL PARA PRECIFICACAO
=====================================================================

Le XMLs de NF-e (somente leitura), extrai 1 linha por ITEM e gera:

  01_INVENTARIO_XML         1 linha por arquivo encontrado
  02_BASE_ITENS_XML         1 linha por item (XML + nota + item)
  02B_TAGS_ITEM_LONGA       (CSV) 1 linha por tag folha de cada item (garantia de completude)
  03_AMOSTRA_FISCAL         amostra representativa com MOTIVO_SELECAO
  04_DIVERGENCIAS_PRODUTOS  mesmo EAN com tratamento diferente entre documentos
  05_MAPA_CENARIOS_FISCAIS  combinacoes fiscais observadas
  05B_GRUPOS_TAGS_OBSERVADOS grupos/tags fiscais realmente presentes nos XMLs
  06_LOG_VALIDACAO          checagens quantitativas/estruturais/amostragem
  06B_LOG_ERROS             erros e avisos por arquivo
  07_DICIONARIO_CAMPOS      dicionario de todas as colunas

REGRAS
  * Os XMLs de origem NUNCA sao alterados, movidos, renomeados ou regravados.
    Arquivos sao abertos apenas em modo leitura binaria ("rb").
  * Identificadores (chave, CNPJ, EAN, NCM, CEST, CFOP, CST, codigos) sao
    gravados como TEXTO.
  * Campo inexistente no XML fica VAZIO (nao vira zero).
  * Colunas CLASS_DOC_* sao classificacao documental (organizacao), nao
    conclusao fiscal. Colunas INDICIO_* sao hipoteses que exigem validacao.
  * Nenhuma regra tributaria e aplicada como verdade.

USO
  python extrator_base_fiscal_xml.py --origem "<pasta BASE_XML>" --saida "<pasta de saida>"
  Opcoes: --excluir "<nome de pasta>" (repetivel)  --amostra-por-cenario 3  --max-eans-divergentes-amostra 150
          --sem-zip  --sem-xlsx

Requisitos: Python 3.9+. openpyxl (opcional, para .xlsx). rarfile + unrar (opcional, para .rar).
"""

import argparse
import csv
import datetime as dt
import hashlib
import io
import json
import os
import re
import sys
import zipfile
import xml.etree.ElementTree as ET
from collections import Counter, OrderedDict, defaultdict

VERSAO_EXTRATOR = "1.0.0"

XLSX_MAX_LINHAS = 1_048_000
XLSX_MAX_CEL = 32_000

# ----------------------------------------------------------------------------
# Utilidades XML
# ----------------------------------------------------------------------------

def local(tag):
    """Remove namespace de uma tag."""
    if not isinstance(tag, str):
        return ""
    return tag.rsplit("}", 1)[-1]


def filho(el, nome):
    if el is None:
        return None
    for c in el:
        if local(c.tag) == nome:
            return c
    return None


def filhos(el, nome):
    if el is None:
        return []
    return [c for c in el if local(c.tag) == nome]


def caminho(el, *nomes):
    atual = el
    for n in nomes:
        atual = filho(atual, n)
        if atual is None:
            return None
    return atual


def txt(el, *nomes):
    """Texto de el/nomes... ; None se ausente (ausente != vazio != zero)."""
    alvo = caminho(el, *nomes) if nomes else el
    if alvo is None:
        return None
    t = alvo.text
    if t is None:
        return ""
    return t.strip()


def folhas(el, prefixo=""):
    """Lista (caminho_relativo, tag, valor) de todas as folhas de el (inclui atributos)."""
    out = []
    if el is None:
        return out
    base = (prefixo + "/" if prefixo else "") + local(el.tag)
    for k, v in el.attrib.items():
        out.append((base + "/@" + local(k), "@" + local(k), v))
    kids = list(el)
    if not kids:
        out.append((base, local(el.tag), (el.text or "").strip()))
    else:
        for k in kids:
            out.extend(folhas(k, base))
    return out


def num(v):
    """Converte texto numerico do XML em float. None se ausente/invalido."""
    if v is None or v == "":
        return None
    try:
        return float(v)
    except ValueError:
        return None


def so_digitos(s):
    return re.sub(r"\D", "", s or "")


def gtin_valido(code):
    """Valida digito verificador GTIN-8/12/13/14."""
    if not code or not code.isdigit() or len(code) not in (8, 12, 13, 14):
        return False
    digits = [int(c) for c in code]
    check = digits[-1]
    body = digits[:-1][::-1]
    total = sum(d * (3 if i % 2 == 0 else 1) for i, d in enumerate(body))
    return (10 - total % 10) % 10 == check


def chave_dv_ok(ch):
    if not ch or len(ch) != 44 or not ch.isdigit():
        return False
    pesos = [2, 3, 4, 5, 6, 7, 8, 9]
    soma = 0
    for i, d in enumerate(reversed(ch[:43])):
        soma += int(d) * pesos[i % 8]
    r = soma % 11
    dv = 0 if r in (0, 1) else 11 - r
    return dv == int(ch[43])


def sha256_bytes(b):
    return hashlib.sha256(b).hexdigest()


def parse_xml_bytes(data):
    """Tenta parsear; retorna (root, aviso). Nunca altera o arquivo de origem."""
    aviso = ""
    try:
        return ET.fromstring(data), aviso
    except ET.ParseError as e1:
        erro1 = str(e1)
    # Fallbacks em memoria (o arquivo original nao e tocado)
    tentativa = data.lstrip(b"\xef\xbb\xbf\r\n\t ")
    if tentativa != data:
        try:
            return ET.fromstring(tentativa), "PARSE_APOS_REMOVER_BOM_OU_ESPACOS_INICIAIS(memoria)"
        except ET.ParseError:
            pass
    try:
        texto = data.decode("latin-1")
        texto = re.sub(r"^\s*<\?xml[^>]*\?>", "", texto)
        return ET.fromstring(texto), "PARSE_COM_FALLBACK_LATIN1(memoria)"
    except Exception:
        pass
    raise ValueError(erro1)


# ----------------------------------------------------------------------------
# Definicao de colunas (registry -> dicionario)
# ----------------------------------------------------------------------------
# tipo: TEXTO | NUMERO | DATA_TEXTO | FLAG
# natureza: ORIGINAL | CONTROLE | CALCULADO | CLASSIF_DOC | INDICIO

COLS = OrderedDict()


def col(nome, desc, origem_tag, tipo="TEXTO", natureza="ORIGINAL", formula=""):
    COLS[nome] = dict(nome=nome, desc=desc, tag=origem_tag, tipo=tipo,
                      natureza=natureza, formula=formula)
    return nome


# --- Rastreabilidade / controle
col("ID_REGISTRO", "Identificador tecnico unico do item: CHAVE_NFE + '-' + NUMERO_ITEM (3 digitos). Nao substitui campos originais.", "infNFe/@Id + det/@nItem", "TEXTO", "CONTROLE", "CHAVE_NFE & '-' & TEXTO(NUMERO_ITEM;'000')")
col("ARQ_CAMINHO_COMPLETO", "Caminho completo do XML de origem (ou caminho do compactado + '::' + arquivo interno).", "sistema de arquivos", "TEXTO", "CONTROLE")
col("ARQ_CAMINHO_RELATIVO", "Caminho relativo a pasta de origem.", "sistema de arquivos", "TEXTO", "CONTROLE")
col("ARQ_NOME", "Nome do arquivo XML.", "sistema de arquivos", "TEXTO", "CONTROLE")
col("ARQ_PASTA_LOJA_CNPJ", "Primeira pasta do caminho com 14 digitos (CNPJ da loja pela organizacao de pastas).", "sistema de arquivos", "TEXTO", "CONTROLE")
col("ARQ_PASTA_TIPO", "Subpasta Compras/Entradas/Saidas quando existir no caminho.", "sistema de arquivos", "TEXTO", "CONTROLE")
col("ARQ_HASH_SHA256", "Hash SHA-256 do arquivo de origem (integridade e duplicidade).", "conteudo binario", "TEXTO", "CONTROLE")
col("QTD_ARQUIVOS_MESMA_CHAVE", "Quantidade de arquivos encontrados com a mesma chave de NF-e (copias/duplicidades).", "inventario", "NUMERO", "CONTROLE")

# --- Cabecalho da nota
col("CHAVE_NFE", "Chave de acesso (44 digitos).", "infNFe/@Id (sem 'NFe') ou protNFe/infProt/chNFe", "TEXTO")
col("MODELO", "Modelo do documento fiscal (55 NF-e, 65 NFC-e).", "ide/mod")
col("SERIE", "Serie.", "ide/serie")
col("NUMERO_NF", "Numero da nota.", "ide/nNF")
col("DATA_EMISSAO", "Data/hora de emissao como no XML.", "ide/dhEmi (ou dEmi)", "DATA_TEXTO")
col("UF_EMITENTE_CODIGO", "Codigo IBGE da UF do emitente.", "ide/cUF")
col("TIPO_NF", "0=entrada, 1=saida (do ponto de vista do emitente).", "ide/tpNF")
col("ID_DEST", "1=interna, 2=interestadual, 3=exterior.", "ide/idDest")
col("FINALIDADE_NFE", "1=normal, 2=complementar, 3=ajuste, 4=devolucao.", "ide/finNFe")
col("IND_FINAL", "Indicador de consumidor final.", "ide/indFinal")
col("NATUREZA_OPERACAO", "Natureza da operacao.", "ide/natOp")
col("EMIT_CNPJ", "CNPJ do emitente (fornecedor).", "emit/CNPJ")
col("EMIT_CPF", "CPF do emitente quando pessoa fisica.", "emit/CPF")
col("EMIT_RAZAO_SOCIAL", "Razao social do emitente.", "emit/xNome")
col("EMIT_FANTASIA", "Nome fantasia do emitente.", "emit/xFant")
col("EMIT_UF", "UF do emitente.", "emit/enderEmit/UF")
col("EMIT_IE", "IE do emitente.", "emit/IE")
col("EMIT_IEST", "IE do substituto tributario.", "emit/IEST")
col("EMIT_CRT", "Codigo de regime tributario: 1=Simples, 2=Simples excesso, 3=Normal, 4=MEI.", "emit/CRT")
col("DEST_CNPJ", "CNPJ do destinatario (loja).", "dest/CNPJ")
col("DEST_CPF", "CPF do destinatario.", "dest/CPF")
col("DEST_RAZAO_SOCIAL", "Razao social do destinatario.", "dest/xNome")
col("DEST_UF", "UF do destinatario.", "dest/enderDest/UF")
col("DEST_IE", "IE do destinatario.", "dest/IE")
col("DEST_IND_IE", "Indicador de IE do destinatario.", "dest/indIEDest")
col("PROT_CSTAT", "Status do protocolo de autorizacao (100=autorizado).", "protNFe/infProt/cStat")
col("PROT_MOTIVO", "Motivo do status do protocolo.", "protNFe/infProt/xMotivo")
col("PROT_NUMERO", "Numero do protocolo.", "protNFe/infProt/nProt")
col("EVENTO_CANCELAMENTO_ENCONTRADO", "S quando foi encontrado XML de evento de cancelamento (110111) para a chave.", "procEventoNFe/evento/infEvento/tpEvento", "FLAG", "CONTROLE")
col("NF_VALOR_TOTAL", "Valor total da nota.", "total/ICMSTot/vNF", "NUMERO")
col("NF_QTD_ITENS", "Quantidade de elementos det na nota.", "contagem de det", "NUMERO", "CONTROLE")

# --- Produto
col("NUMERO_ITEM", "Numero sequencial do item na nota.", "det/@nItem", "TEXTO")
col("COD_PRODUTO_FORNECEDOR", "Codigo do produto no fornecedor (cProd). Nao e identificador global.", "det/prod/cProd")
col("EAN", "GTIN comercial (cEAN) exatamente como no XML.", "det/prod/cEAN")
col("EAN_TRIB", "GTIN tributavel (cEANTrib) exatamente como no XML.", "det/prod/cEANTrib")
col("DESCRICAO_PRODUTO", "Descricao do produto no XML.", "det/prod/xProd")
col("NCM", "NCM.", "det/prod/NCM")
col("CEST", "CEST.", "det/prod/CEST")
col("IND_ESCALA", "Indicador de escala relevante.", "det/prod/indEscala")
col("CNPJ_FABRICANTE", "CNPJ do fabricante (quando indEscala=N).", "det/prod/CNPJFab")
col("COD_BENEF_FISCAL", "Codigo de beneficio fiscal (cBenef).", "det/prod/cBenef")
col("EXTIPI", "EX da TIPI.", "det/prod/EXTIPI")
col("CFOP", "CFOP do item.", "det/prod/CFOP")
col("UNID_COMERCIAL", "Unidade comercial.", "det/prod/uCom")
col("QTD_COMERCIAL", "Quantidade comercial.", "det/prod/qCom", "NUMERO")
col("VALOR_UNIT_COMERCIAL", "Valor unitario comercial.", "det/prod/vUnCom", "NUMERO")
col("VALOR_PRODUTO", "Valor bruto do produto (vProd).", "det/prod/vProd", "NUMERO")
col("UNID_TRIBUTAVEL", "Unidade tributavel.", "det/prod/uTrib")
col("QTD_TRIBUTAVEL", "Quantidade tributavel.", "det/prod/qTrib", "NUMERO")
col("VALOR_UNIT_TRIBUTAVEL", "Valor unitario tributavel.", "det/prod/vUnTrib", "NUMERO")
col("VALOR_FRETE", "Frete do item.", "det/prod/vFrete", "NUMERO")
col("VALOR_SEGURO", "Seguro do item.", "det/prod/vSeg", "NUMERO")
col("VALOR_DESCONTO", "Desconto do item.", "det/prod/vDesc", "NUMERO")
col("VALOR_OUTRAS_DESP", "Outras despesas acessorias do item.", "det/prod/vOutro", "NUMERO")
col("IND_TOTAL", "Indica se vProd compoe o total da NF.", "det/prod/indTot")
col("PEDIDO_COMPRA", "Numero do pedido de compra (xPed).", "det/prod/xPed")
col("ITEM_PEDIDO", "Item do pedido (nItemPed).", "det/prod/nItemPed")
col("NUMERO_FCI", "Numero da FCI.", "det/prod/nFCI")
col("ANVISA_CODIGO", "Codigo ANVISA do medicamento.", "det/prod/med/cProdANVISA")
col("ANVISA_MOTIVO_ISENCAO", "Motivo de isencao do registro ANVISA.", "det/prod/med/xMotivoIsencao")
col("PMC_XML", "Preco maximo ao consumidor informado no XML (vPMC).", "det/prod/med/vPMC", "NUMERO")
col("RASTRO_QTD_LOTES", "Quantidade de grupos rastro (lotes) no item.", "det/prod/rastro", "NUMERO", "CONTROLE")
col("RASTRO_LOTES", "Lotes no formato nLote|qLote|dFab|dVal separados por ' / '.", "det/prod/rastro/*")
col("VALOR_TOTAL_ITEM_XML", "Valor total do item (vItem) quando informado (layout novo).", "det/vItem", "NUMERO")
col("VALOR_TOT_TRIB_ITEM", "Valor aproximado total de tributos (Lei 12.741).", "det/imposto/vTotTrib", "NUMERO")
col("INF_AD_PROD", "Informacoes adicionais do item.", "det/infAdProd")

# --- ICMS
col("ICMS_GRUPO_TAG", "Tag do grupo de ICMS encontrado (ICMS00, ICMS10, ICMS60, ICMSSN102, ...).", "det/imposto/ICMS/*")
col("ICMS_ORIGEM", "Origem da mercadoria.", "ICMSxx/orig")
col("ICMS_CST", "CST do ICMS.", "ICMSxx/CST")
col("ICMS_CSOSN", "CSOSN (Simples Nacional).", "ICMSSNxxx/CSOSN")
_ICMS_CAMPOS = [
    ("modBC", "Modalidade da BC do ICMS", "TEXTO"), ("pRedBC", "% reducao da BC", "NUMERO"),
    ("vBC", "BC do ICMS", "NUMERO"), ("pICMS", "Aliquota ICMS", "NUMERO"), ("vICMSOp", "ICMS da operacao (diferimento)", "NUMERO"),
    ("pDif", "% diferimento", "NUMERO"), ("vICMSDif", "ICMS diferido", "NUMERO"), ("vICMS", "Valor ICMS", "NUMERO"),
    ("vBCFCP", "BC FCP", "NUMERO"), ("pFCP", "% FCP", "NUMERO"), ("vFCP", "Valor FCP", "NUMERO"),
    ("pFCPDif", "% FCP diferido", "NUMERO"), ("vFCPDif", "FCP diferido", "NUMERO"), ("vFCPEfet", "FCP efetivo", "NUMERO"),
    ("modBCST", "Modalidade BC ST", "TEXTO"), ("pMVAST", "MVA ST", "NUMERO"), ("pRedBCST", "% reducao BC ST", "NUMERO"),
    ("vBCST", "BC ICMS ST", "NUMERO"), ("pICMSST", "Aliquota ICMS ST", "NUMERO"), ("vICMSST", "Valor ICMS ST", "NUMERO"),
    ("vBCFCPST", "BC FCP ST", "NUMERO"), ("pFCPST", "% FCP ST", "NUMERO"), ("vFCPST", "Valor FCP ST", "NUMERO"),
    ("vICMSSTDeson", "ICMS ST desonerado", "NUMERO"), ("motDesICMSST", "Motivo desoneracao ICMS ST", "TEXTO"),
    ("vBCSTRet", "BC ICMS ST retido anteriormente", "NUMERO"), ("pST", "Aliquota suportada pelo consumidor final (ST retido)", "NUMERO"),
    ("vICMSSubstituto", "ICMS proprio do substituto", "NUMERO"), ("vICMSSTRet", "ICMS ST retido anteriormente", "NUMERO"),
    ("vBCFCPSTRet", "BC FCP ST retido", "NUMERO"), ("pFCPSTRet", "% FCP ST retido", "NUMERO"), ("vFCPSTRet", "FCP ST retido", "NUMERO"),
    ("pRedBCEfet", "% reducao BC efetiva", "NUMERO"), ("vBCEfet", "BC efetiva", "NUMERO"), ("pICMSEfet", "Aliquota ICMS efetiva", "NUMERO"),
    ("vICMSEfet", "ICMS efetivo", "NUMERO"), ("vICMSDeson", "ICMS desonerado", "NUMERO"), ("motDesICMS", "Motivo desoneracao ICMS", "TEXTO"),
    ("indDeduzDeson", "Indica deducao do desonerado", "TEXTO"), ("pCredSN", "% credito Simples Nacional", "NUMERO"),
    ("vCredICMSSN", "Credito ICMS Simples Nacional", "NUMERO"), ("vBCSTDest", "BC ST UF destino (repasse)", "NUMERO"),
    ("vICMSSTDest", "ICMS ST UF destino (repasse)", "NUMERO"), ("pBCOp", "% BC operacao propria (ICMSPart)", "NUMERO"),
    ("UFST", "UF ST (ICMSPart)", "TEXTO"), ("qBCMono", "Qtd BC monofasico", "NUMERO"), ("adRemICMS", "Ad rem ICMS", "NUMERO"),
    ("vICMSMono", "ICMS monofasico", "NUMERO"), ("vICMSMonoOp", "ICMS monofasico operacao", "NUMERO"), ("vICMSMonoDif", "ICMS monofasico diferido", "NUMERO"),
    ("qBCMonoReten", "Qtd BC mono retencao", "NUMERO"), ("adRemICMSReten", "Ad rem retencao", "NUMERO"), ("vICMSMonoReten", "ICMS mono retencao", "NUMERO"),
    ("qBCMonoRet", "Qtd BC mono retido anteriormente", "NUMERO"), ("adRemICMSRet", "Ad rem retido anteriormente", "NUMERO"),
    ("vICMSMonoRet", "ICMS mono retido anteriormente", "NUMERO"),
]
for t, d, tp in _ICMS_CAMPOS:
    col("ICMS_" + t, d + ".", "det/imposto/ICMS/ICMSxx/" + t, tp)
col("ICMS_TAGS_NAO_MAPEADAS", "Tags do grupo ICMS sem coluna propria (tag=valor). Garantia de nao perda.", "det/imposto/ICMS/ICMSxx/*")

# DIFAL
for t, d in [("vBCUFDest", "BC UF destino"), ("vBCFCPUFDest", "BC FCP UF destino"), ("pFCPUFDest", "% FCP UF destino"),
             ("pICMSUFDest", "Aliquota interna UF destino"), ("pICMSInter", "Aliquota interestadual"),
             ("pICMSInterPart", "% partilha"), ("vFCPUFDest", "FCP UF destino"), ("vICMSUFDest", "ICMS UF destino"),
             ("vICMSUFRemet", "ICMS UF remetente")]:
    col("DIFAL_" + t, d + ".", "det/imposto/ICMSUFDest/" + t, "NUMERO")

# --- IPI
col("IPI_GRUPO_TAG", "Tag do subgrupo IPI (IPITrib / IPINT). Vazio se nao ha grupo IPI.", "det/imposto/IPI/*")
col("IPI_CST", "CST do IPI.", "IPI/IPITrib|IPINT/CST")
col("IPI_COD_ENQ", "Codigo de enquadramento do IPI.", "IPI/cEnq")
col("IPI_CNPJ_PRODUTOR", "CNPJ do produtor.", "IPI/CNPJProd")
col("IPI_COD_SELO", "Codigo do selo.", "IPI/cSelo")
col("IPI_QTD_SELO", "Quantidade de selos.", "IPI/qSelo", "NUMERO")
for t, d in [("vBC", "BC IPI"), ("pIPI", "Aliquota IPI"), ("qUnid", "Quantidade (pauta)"), ("vUnid", "Valor por unidade"), ("vIPI", "Valor IPI")]:
    col("IPI_" + t, d + ".", "IPI/IPITrib/" + t, "NUMERO")
col("IPI_DEVOL_PERC", "% mercadoria devolvida.", "det/impostoDevol/pDevol", "NUMERO")
col("IPI_DEVOL_VALOR", "IPI devolvido.", "det/impostoDevol/IPI/vIPIDevol", "NUMERO")

# --- II
for t, d in [("vBC", "BC II"), ("vDespAdu", "Despesas aduaneiras"), ("vII", "Valor II"), ("vIOF", "Valor IOF")]:
    col("II_" + t, d + ".", "det/imposto/II/" + t, "NUMERO")

# --- PIS / COFINS
for trib in ("PIS", "COFINS"):
    col(trib + "_GRUPO_TAG", "Tag do grupo %s encontrado (%sAliq, %sQtde, %sNT, %sOutr)." % ((trib,) * 5), "det/imposto/%s/*" % trib)
    col(trib + "_CST", "CST %s." % trib, "%sxx/CST" % trib)
    col(trib + "_VBC", "BC %s." % trib, "%sxx/vBC" % trib, "NUMERO")
    col(trib + "_ALIQ_PERC", "Aliquota percentual %s." % trib, "%sxx/p%s" % (trib, trib), "NUMERO")
    col(trib + "_QTD_BC_PROD", "Quantidade vendida (aliquota por unidade).", "%sxx/qBCProd" % trib, "NUMERO")
    col(trib + "_ALIQ_VALOR", "Aliquota em reais por unidade.", "%sxx/vAliqProd" % trib, "NUMERO")
    col(trib + "_VALOR", "Valor %s." % trib, "%sxx/v%s" % (trib, trib), "NUMERO")
    col(trib + "ST_VALOR", "Valor %s ST (grupo %sST)." % (trib, trib), "det/imposto/%sST/v%s" % (trib, trib), "NUMERO")
    col(trib + "_TAGS_NAO_MAPEADAS", "Tags do grupo %s sem coluna propria (tag=valor)." % trib, "det/imposto/%s/*" % trib)

# --- Reforma tributaria IBS/CBS/IS
col("IBSCBS_PRESENTE", "S se o grupo IBSCBS existe no item.", "det/imposto/IBSCBS", "FLAG", "CONTROLE")
col("IBSCBS_CST", "CST IBS/CBS.", "IBSCBS/CST")
col("IBSCBS_CCLASSTRIB", "Codigo de classificacao tributaria IBS/CBS.", "IBSCBS/cClassTrib")
col("IBSCBS_VBC", "BC IBS/CBS.", "IBSCBS/gIBSCBS/vBC", "NUMERO")
col("IBS_UF_ALIQ", "Aliquota IBS UF.", "gIBSUF/pIBSUF", "NUMERO")
col("IBS_UF_RED_ALIQ", "% reducao aliquota IBS UF.", "gIBSUF/gRed/pRedAliq", "NUMERO")
col("IBS_UF_ALIQ_EFET", "Aliquota efetiva IBS UF.", "gIBSUF/gRed/pAliqEfet", "NUMERO")
col("IBS_UF_VALOR", "Valor IBS UF.", "gIBSUF/vIBSUF", "NUMERO")
col("IBS_MUN_ALIQ", "Aliquota IBS municipal.", "gIBSMun/pIBSMun", "NUMERO")
col("IBS_MUN_RED_ALIQ", "% reducao aliquota IBS municipal.", "gIBSMun/gRed/pRedAliq", "NUMERO")
col("IBS_MUN_ALIQ_EFET", "Aliquota efetiva IBS municipal.", "gIBSMun/gRed/pAliqEfet", "NUMERO")
col("IBS_MUN_VALOR", "Valor IBS municipal.", "gIBSMun/vIBSMun", "NUMERO")
col("IBS_VALOR", "Valor IBS total do item.", "gIBSCBS/vIBS", "NUMERO")
col("CBS_ALIQ", "Aliquota CBS.", "gCBS/pCBS", "NUMERO")
col("CBS_RED_ALIQ", "% reducao aliquota CBS.", "gCBS/gRed/pRedAliq", "NUMERO")
col("CBS_ALIQ_EFET", "Aliquota efetiva CBS.", "gCBS/gRed/pAliqEfet", "NUMERO")
col("CBS_VALOR", "Valor CBS.", "gCBS/vCBS", "NUMERO")
col("IBSCBS_TAGS_NAO_MAPEADAS", "Demais tags do grupo IBSCBS (tag=valor).", "det/imposto/IBSCBS/*")
col("IS_PRESENTE", "S se grupo IS (Imposto Seletivo) existe.", "det/imposto/IS", "FLAG", "CONTROLE")
col("IS_TAGS", "Tags do grupo IS (tag=valor).", "det/imposto/IS/*")
col("OUTROS_GRUPOS_IMPOSTO", "Outros grupos dentro de imposto sem tratamento proprio (nome=tags).", "det/imposto/*")
col("PROD_TAGS_NAO_MAPEADAS", "Tags de det/prod sem coluna propria (tag=valor).", "det/prod/*")

# --- Totais e informacoes da nota
for t, d in [("vBC", "BC ICMS"), ("vICMS", "ICMS"), ("vICMSDeson", "ICMS desonerado"), ("vFCP", "FCP"), ("vBCST", "BC ST"),
             ("vST", "ICMS ST"), ("vFCPST", "FCP ST"), ("vFCPSTRet", "FCP ST retido"), ("vProd", "Produtos"),
             ("vFrete", "Frete"), ("vSeg", "Seguro"), ("vDesc", "Desconto"), ("vII", "II"), ("vIPI", "IPI"),
             ("vIPIDevol", "IPI devolvido"), ("vPIS", "PIS"), ("vCOFINS", "COFINS"), ("vOutro", "Outras despesas")]:
    col("NF_TOT_" + t, "Total da nota - " + d + ".", "total/ICMSTot/" + t, "NUMERO")
col("NF_MOD_FRETE", "Modalidade do frete.", "transp/modFrete")
col("NF_INF_AD_FISCO", "Informacoes adicionais de interesse do fisco.", "infAdic/infAdFisco")
col("NF_INF_COMPLEMENTAR", "Informacoes complementares.", "infAdic/infCpl")

# --- Calculado / validacao (nao substitui original)
col("CALC_DIF_VPROD_QCOM_X_VUNCOM", "Diferenca VALOR_PRODUTO - QTD_COMERCIAL*VALOR_UNIT_COMERCIAL. Checagem de consistencia.", "calculado", "NUMERO", "CALCULADO", "VALOR_PRODUTO - QTD_COMERCIAL * VALOR_UNIT_COMERCIAL")
col("CALC_DIF_VICMS_BC_X_ALIQ", "Diferenca ICMS_vICMS - ICMS_vBC*ICMS_pICMS/100. Checagem de consistencia.", "calculado", "NUMERO", "CALCULADO", "ICMS_vICMS - ICMS_vBC * ICMS_pICMS / 100")
col("EAN_GTIN_VALIDO", "S se EAN tem 8/12/13/14 digitos e digito verificador correto.", "calculado sobre cEAN", "FLAG", "CALCULADO", "regra GS1 de digito verificador")
col("EAN_CHAVE_VINCULO", "EAN usado para vinculo entre XMLs: cEAN se valido, senao cEANTrib se valido, senao vazio.", "calculado sobre cEAN/cEANTrib", "TEXTO", "CALCULADO", "SE(EAN valido; EAN; SE(EAN_TRIB valido; EAN_TRIB; ''))")

# --- Classificacao documental
col("CLASS_DOC_TIPO_OPERACAO", "Organizacao da operacao: MESMA RAIZ CNPJ (provavel transferencia interna) x TERCEIRO; finalidade.", "emit/CNPJ, dest/CNPJ, finNFe, tpNF", "TEXTO", "CLASSIF_DOC")
col("CLASS_DOC_ICMS", "Descricao documental do tratamento de ICMS conforme grupo/CST no XML.", "ICMS grupo/CST/CSOSN/valores", "TEXTO", "CLASSIF_DOC")
col("FLAG_ST_DESTACADA", "S se vICMSST > 0 no item.", "ICMSxx/vICMSST", "FLAG", "CLASSIF_DOC")
col("FLAG_BC_ST_INFORMADA", "S se vBCST > 0 no item.", "ICMSxx/vBCST", "FLAG", "CLASSIF_DOC")
col("FLAG_ST_RETIDA_ANTERIOR_INFO", "S se vBCSTRet ou vICMSSTRet > 0.", "ICMSxx/vBCSTRet, vICMSSTRet", "FLAG", "CLASSIF_DOC")
col("FLAG_ICMS_DESONERADO", "S se vICMSDeson > 0.", "ICMSxx/vICMSDeson", "FLAG", "CLASSIF_DOC")
col("FLAG_DIFERIMENTO", "S se pDif ou vICMSDif informado > 0.", "ICMSxx/pDif, vICMSDif", "FLAG", "CLASSIF_DOC")
col("FLAG_FCP", "S se algum FCP > 0 (vFCP, vFCPST, vFCPSTRet).", "ICMSxx/vFCP*", "FLAG", "CLASSIF_DOC")
col("FLAG_IPI_VALOR", "S se vIPI > 0.", "IPI/IPITrib/vIPI", "FLAG", "CLASSIF_DOC")
col("FLAG_DIFAL", "S se grupo ICMSUFDest existe.", "det/imposto/ICMSUFDest", "FLAG", "CLASSIF_DOC")
col("CLASS_DOC_PIS", "Descricao documental do PIS conforme grupo/CST/valores no XML.", "PIS grupo/CST/valores", "TEXTO", "CLASSIF_DOC")
col("CLASS_DOC_COFINS", "Descricao documental da COFINS conforme grupo/CST/valores no XML.", "COFINS grupo/CST/valores", "TEXTO", "CLASSIF_DOC")
col("CLASS_DOC_IPI", "Descricao documental do IPI.", "IPI grupo/CST/valores", "TEXTO", "CLASSIF_DOC")
col("CENARIO_FISCAL_ID", "Codigo do cenario fiscal documental (ver 05_MAPA_CENARIOS_FISCAIS).", "combinacao de grupos/CST", "TEXTO", "CLASSIF_DOC")
col("CENARIO_FISCAL_CHAVE", "Combinacao: ICMS grupo|CST/CSOSN|ST|PIS grupo/CST|COFINS grupo/CST|IPI grupo/CST|IBSCBS CST.", "combinacao de grupos/CST", "TEXTO", "CLASSIF_DOC")
col("ALERTAS_DOCUMENTAIS", "Inconsistencias/atipicidades objetivas detectadas no documento (separadas por ' | ').", "regras de consistencia documentadas", "TEXTO", "CLASSIF_DOC")
col("EAN_COM_DIVERGENCIA", "S se o EAN_CHAVE_VINCULO aparece com tratamento diferente em outro registro (ver 04).", "04_DIVERGENCIAS_PRODUTOS", "FLAG", "CLASSIF_DOC")
col("INDICIO_PARA_VALIDACAO", "HIPOTESE derivada do XML que exige validacao fiscal/legal. Nao e fato.", "derivado", "TEXTO", "INDICIO")
col("STATUS_VALIDACAO_FISCAL", "Sempre PENDENTE nesta fase. Sera preenchido apos validacao fiscal/contabil.", "-", "TEXTO", "INDICIO")

ITEM_COLS = list(COLS.keys())
IDX = {c: i for i, c in enumerate(ITEM_COLS)}


# textos que se repetem muito entre itens: uma unica copia em memoria
INTERNAR = ["CLASS_DOC_TIPO_OPERACAO", "CLASS_DOC_ICMS", "CLASS_DOC_PIS", "CLASS_DOC_COFINS", "CLASS_DOC_IPI",
            "CENARIO_FISCAL_CHAVE", "ALERTAS_DOCUMENTAIS", "INDICIO_PARA_VALIDACAO", "NCM", "CEST", "CFOP",
            "UNID_COMERCIAL", "UNID_TRIBUTAVEL", "ICMS_GRUPO_TAG", "ICMS_CST", "ICMS_ORIGEM", "PIS_GRUPO_TAG", "PIS_CST",
            "COFINS_GRUPO_TAG", "COFINS_CST", "IPI_GRUPO_TAG", "IPI_CST", "IBSCBS_CST", "IBSCBS_CCLASSTRIB", "EAN", "EAN_TRIB"]


class Row(object):
    """Linha compacta de item (lista em vez de dict) para suportar volumes grandes."""
    __slots__ = ("v",)

    def __init__(self, d):
        self.v = [d.get(c) for c in ITEM_COLS]

    def __getitem__(self, k):
        return self.v[IDX[k]]

    def __setitem__(self, k, val):
        self.v[IDX[k]] = val

    def get(self, k, default=None):
        i = IDX.get(k)
        return self.v[i] if i is not None else default

# Colunas exclusivas da amostra
AMOSTRA_EXTRA = OrderedDict([
    ("AMOSTRA_SEQ", "Sequencial da amostra."),
    ("MOTIVO_SELECAO", "Motivo(s) pelo(s) qual(is) o registro entrou na amostra."),
])

# Colunas de inventario (compativeis com CONTROLE_XML_MASTER/ARQUIVOS + extras)
INV_COLS = [
    ("CAMINHO_COMPLETO", "Caminho completo do arquivo (compactado::interno quando aplicavel)."),
    ("CAMINHO_RELATIVO", "Caminho relativo a origem."),
    ("NOME_ARQUIVO", "Nome do arquivo."),
    ("EXTENSAO", "Extensao."),
    ("TAMANHO_BYTES", "Tamanho em bytes."),
    ("DATA_MODIFICACAO", "Data de modificacao do arquivo (sistema)."),
    ("HASH_SHA256", "SHA-256 do conteudo."),
    ("TIPO_DOCUMENTO", "nfeProc / NFe_SEM_PROTOCOLO / procEventoNFe / resNFe / CTe / XML_OUTRO / NAO_XML / COMPACTADO."),
    ("MODELO", "ide/mod."),
    ("CHAVE_NFE", "Chave de acesso."),
    ("CHAVE_DV_OK", "S se digito verificador da chave confere."),
    ("CHAVE_IGUAL_NOME_ARQUIVO", "S se a chave aparece no nome do arquivo; N se o nome tem outra chave; vazio se nome sem chave."),
    ("NUMERO_NF", "ide/nNF."),
    ("SERIE", "ide/serie."),
    ("DATA_EMISSAO", "ide/dhEmi."),
    ("EMIT_CNPJ", "emit/CNPJ."),
    ("EMIT_RAZAO_SOCIAL", "emit/xNome."),
    ("DEST_CNPJ", "dest/CNPJ."),
    ("DEST_RAZAO_SOCIAL", "dest/xNome."),
    ("NATUREZA_OPERACAO", "ide/natOp."),
    ("QTD_ITENS", "Quantidade de det."),
    ("QTD_ITENS_EXTRAIDOS", "Itens efetivamente gravados na base (0 se duplicado)."),
    ("PROT_CSTAT", "protNFe/infProt/cStat."),
    ("EVENTO_TIPO", "Tipo de evento (procEventoNFe)."),
    ("STATUS", "LIDO_OK / LIDO_COM_AVISO / DUPLICADO_HASH / DUPLICADO_CHAVE / ERRO_LEITURA / NAO_NFE / EVENTO / COMPACTADO_PROCESSADO / COMPACTADO_NAO_PROCESSADO / IGNORADO_NAO_XML."),
    ("DUPLICADO_DE", "Caminho do arquivo canonico quando duplicado."),
    ("ERRO", "Mensagem de erro/aviso."),
    ("DATA_PROCESSAMENTO", "Data/hora da execucao."),
]


# ----------------------------------------------------------------------------
# Extracao
# ----------------------------------------------------------------------------

class Log:
    def __init__(self):
        self.erros = []

    def add(self, nivel, etapa, arquivo, erro, causa, impacto, acao):
        self.erros.append(OrderedDict([
            ("NIVEL", nivel), ("ETAPA", etapa), ("ARQUIVO", arquivo), ("ERRO", erro),
            ("CAUSA_PROVAVEL", causa), ("IMPACTO", impacto), ("ACAO_TOMADA", acao)]))


def pares_txt(lista):
    return " ; ".join("%s=%s" % (k, v) for k, v in lista)


def classificar_documento(root):
    nome = local(root.tag)
    if nome == "nfeProc":
        return "nfeProc"
    if nome == "NFe":
        return "NFe_SEM_PROTOCOLO"
    if nome in ("procEventoNFe", "evento", "retEnvEvento"):
        return "procEventoNFe"
    if nome in ("resNFe", "resEvento"):
        return nome
    if nome in ("cteProc", "CTe"):
        return "CTe"
    if nome in ("nfseProc", "CompNfse", "Nfse"):
        return "NFSe"
    return "XML_OUTRO(" + nome + ")"


def info_pasta(rel):
    partes = re.split(r"[\\/]", rel)
    loja = ""
    tipo = ""
    for p in partes[:-1]:
        if re.fullmatch(r"\d{14}", p) and not loja:
            loja = p
        if p.lower() in ("compras", "entradas", "saidas", "saídas"):
            tipo = p
    return loja, tipo


def fmt_flag(b):
    return "S" if b else "N"


def gt0(v):
    return v is not None and v > 0


ICMS_MAP = {t for t, _, _ in _ICMS_CAMPOS} | {"orig", "CST", "CSOSN"}
PIS_MAP = {"CST", "vBC", "pPIS", "qBCProd", "vAliqProd", "vPIS"}
COFINS_MAP = {"CST", "vBC", "pCOFINS", "qBCProd", "vAliqProd", "vCOFINS"}
PROD_MAP = {"cProd", "cEAN", "cEANTrib", "xProd", "NCM", "CEST", "indEscala", "CNPJFab", "cBenef", "EXTIPI", "CFOP",
            "uCom", "qCom", "vUnCom", "vProd", "uTrib", "qTrib", "vUnTrib", "vFrete", "vSeg", "vDesc", "vOutro",
            "indTot", "xPed", "nItemPed", "nFCI", "med", "rastro"}


def descreve_icms(grupo, cst, csosn, v):
    if not grupo:
        return "SEM GRUPO ICMS NO XML"
    if csosn:
        base = "SIMPLES NACIONAL CSOSN %s (%s)" % (csosn, grupo)
        if gt0(v("vICMSST")):
            base += " COM ST DESTACADA"
        if gt0(v("vICMSSTRet")) or gt0(v("vBCSTRet")):
            base += " COM ST RETIDA INFORMADA"
        return base
    desc = {
        "00": "ICMS TRIBUTADO INTEGRALMENTE NO XML",
        "02": "ICMS MONOFASICO PROPRIO (COMBUSTIVEIS) NO XML",
        "10": "ICMS PROPRIO + GRUPO ST NO XML",
        "15": "ICMS MONOFASICO COM RETENCAO NO XML",
        "20": "ICMS COM REDUCAO DE BC NO XML",
        "30": "ICMS ISENTO/NAO TRIBUTADO + GRUPO ST NO XML",
        "40": "ICMS CST 40 (ISENTA) NO XML",
        "41": "ICMS CST 41 (NAO TRIBUTADA) NO XML",
        "50": "ICMS CST 50 (SUSPENSAO) NO XML",
        "51": "ICMS CST 51 (DIFERIMENTO) NO XML",
        "53": "ICMS MONOFASICO DIFERIDO NO XML",
        "60": "ICMS CST 60 (ST COBRADO ANTERIORMENTE) NO XML",
        "61": "ICMS MONOFASICO RETIDO ANTERIORMENTE NO XML",
        "70": "ICMS COM REDUCAO DE BC + GRUPO ST NO XML",
        "90": "ICMS CST 90 (OUTRAS) NO XML",
    }.get(cst, "ICMS CST %s (%s) NO XML" % (cst, grupo))
    if grupo == "ICMSPart":
        desc = "ICMS PARTILHA (ICMSPart) CST %s NO XML" % cst
    if grupo == "ICMSST":
        desc = "ICMS ST REPASSE (ICMSST) CST %s NO XML" % cst
    if cst in ("10", "30", "70", "90") or grupo in ("ICMSPart",):
        if gt0(v("vICMSST")):
            desc += " - ST DESTACADA (vICMSST>0)"
        elif gt0(v("vBCST")):
            desc += " - BC ST INFORMADA COM vICMSST ZERO"
        elif v("vBCST") is not None:
            desc += " - GRUPO ST SEM VALOR"
    if cst == "60" or grupo == "ICMSST":
        if gt0(v("vICMSSTRet")) or gt0(v("vBCSTRet")):
            desc += " - COM VALORES DE ST RETIDO INFORMADOS"
        else:
            desc += " - SEM VALORES DE ST RETIDO (ZERADOS OU AUSENTES)"
    if gt0(v("vICMSDeson")):
        desc += " - COM ICMS DESONERADO"
    return desc


CST_PISCOFINS_DESC = {
    "01": "OPERACAO TRIBUTAVEL ALIQUOTA BASICA",
    "02": "OPERACAO TRIBUTAVEL ALIQUOTA DIFERENCIADA",
    "03": "OPERACAO TRIBUTAVEL POR QUANTIDADE",
    "04": "OPERACAO TRIBUTAVEL MONOFASICA - REVENDA A ALIQUOTA ZERO (CST 04)",
    "05": "OPERACAO TRIBUTAVEL POR SUBSTITUICAO TRIBUTARIA (CST 05)",
    "06": "OPERACAO TRIBUTAVEL A ALIQUOTA ZERO (CST 06)",
    "07": "OPERACAO ISENTA (CST 07)",
    "08": "OPERACAO SEM INCIDENCIA (CST 08)",
    "09": "OPERACAO COM SUSPENSAO (CST 09)",
    "49": "OUTRAS OPERACOES DE SAIDA (CST 49)",
    "99": "OUTRAS OPERACOES (CST 99)",
}


def descreve_piscofins(trib, grupo, cst, valor, aliq):
    if not grupo:
        return "SEM GRUPO %s NO XML" % trib
    base = "%s GRUPO %s CST %s [DESCRICAO DA TABELA CST PARA O CODIGO INFORMADO PELO EMITENTE: %s]" % (
        trib, grupo, cst or "?", CST_PISCOFINS_DESC.get(cst, "CODIGO SEM DESCRICAO CADASTRADA NO EXTRATOR"))
    if gt0(valor):
        base += " - VALOR DESTACADO"
    else:
        base += " - SEM VALOR DESTACADO"
    if aliq is not None and aliq == 0 and grupo.endswith("Aliq"):
        base += " (ALIQUOTA 0 NO GRUPO Aliq)"
    return base


def descreve_ipi(grupo, cst, vipi):
    if not grupo:
        return "SEM GRUPO IPI NO XML"
    base = "IPI %s CST %s" % (grupo, cst or "?")
    base += " - VALOR DESTACADO" if gt0(vipi) else " - SEM VALOR DESTACADO"
    return base


def extrair_nfe(root, meta, log, eventos_cancel, sink_longa):
    """Retorna lista de Row (1 por item). Tags folha vao para sink_longa (streaming)."""
    nfe = root if local(root.tag) == "NFe" else filho(root, "NFe")
    prot = caminho(root, "protNFe", "infProt")
    inf = filho(nfe, "infNFe")
    if inf is None:
        raise ValueError("infNFe nao encontrado")
    ide = filho(inf, "ide")
    emit = filho(inf, "emit")
    dest = filho(inf, "dest")
    tot = caminho(inf, "total", "ICMSTot")
    chave = (inf.attrib.get("Id") or "").replace("NFe", "")
    if not chave and prot is not None:
        chave = txt(prot, "chNFe") or ""
    dets = filhos(inf, "det")

    cab = OrderedDict()
    cab["CHAVE_NFE"] = chave
    cab["MODELO"] = txt(ide, "mod")
    cab["SERIE"] = txt(ide, "serie")
    cab["NUMERO_NF"] = txt(ide, "nNF")
    cab["DATA_EMISSAO"] = txt(ide, "dhEmi") or txt(ide, "dEmi")
    cab["UF_EMITENTE_CODIGO"] = txt(ide, "cUF")
    cab["TIPO_NF"] = txt(ide, "tpNF")
    cab["ID_DEST"] = txt(ide, "idDest")
    cab["FINALIDADE_NFE"] = txt(ide, "finNFe")
    cab["IND_FINAL"] = txt(ide, "indFinal")
    cab["NATUREZA_OPERACAO"] = txt(ide, "natOp")
    cab["EMIT_CNPJ"] = txt(emit, "CNPJ")
    cab["EMIT_CPF"] = txt(emit, "CPF")
    cab["EMIT_RAZAO_SOCIAL"] = txt(emit, "xNome")
    cab["EMIT_FANTASIA"] = txt(emit, "xFant")
    cab["EMIT_UF"] = txt(emit, "enderEmit", "UF")
    cab["EMIT_IE"] = txt(emit, "IE")
    cab["EMIT_IEST"] = txt(emit, "IEST")
    cab["EMIT_CRT"] = txt(emit, "CRT")
    cab["DEST_CNPJ"] = txt(dest, "CNPJ")
    cab["DEST_CPF"] = txt(dest, "CPF")
    cab["DEST_RAZAO_SOCIAL"] = txt(dest, "xNome")
    cab["DEST_UF"] = txt(dest, "enderDest", "UF")
    cab["DEST_IE"] = txt(dest, "IE")
    cab["DEST_IND_IE"] = txt(dest, "indIEDest")
    cab["PROT_CSTAT"] = txt(prot, "cStat") if prot is not None else None
    cab["PROT_MOTIVO"] = txt(prot, "xMotivo") if prot is not None else None
    cab["PROT_NUMERO"] = txt(prot, "nProt") if prot is not None else None
    cab["EVENTO_CANCELAMENTO_ENCONTRADO"] = "S" if chave in eventos_cancel else "N"
    cab["NF_VALOR_TOTAL"] = num(txt(tot, "vNF"))
    cab["NF_QTD_ITENS"] = len(dets)
    for t in ("vBC", "vICMS", "vICMSDeson", "vFCP", "vBCST", "vST", "vFCPST", "vFCPSTRet", "vProd", "vFrete", "vSeg",
              "vDesc", "vII", "vIPI", "vIPIDevol", "vPIS", "vCOFINS", "vOutro"):
        cab["NF_TOT_" + t] = num(txt(tot, t))
    cab["NF_MOD_FRETE"] = txt(inf, "transp", "modFrete")
    cab["NF_INF_AD_FISCO"] = txt(inf, "infAdic", "infAdFisco")
    cab["NF_INF_COMPLEMENTAR"] = txt(inf, "infAdic", "infCpl")

    # Classificacao de operacao (documental)
    e_raiz = so_digitos(cab["EMIT_CNPJ"])[:8]
    d_raiz = so_digitos(cab["DEST_CNPJ"])[:8]
    if e_raiz and d_raiz and e_raiz == d_raiz:
        op = "MESMA RAIZ CNPJ EMITENTE/DESTINATARIO (PROVAVEL TRANSFERENCIA INTERNA)"
    elif e_raiz:
        op = "EMITENTE TERCEIRO (PROVAVEL COMPRA/ENTRADA DE FORNECEDOR)"
    else:
        op = "EMITENTE SEM CNPJ"
    fin = {"1": "NORMAL", "2": "COMPLEMENTAR", "3": "AJUSTE", "4": "DEVOLUCAO"}.get(cab["FINALIDADE_NFE"] or "", "FINALIDADE " + str(cab["FINALIDADE_NFE"]))
    tp = {"0": "TPNF ENTRADA", "1": "TPNF SAIDA"}.get(cab["TIPO_NF"] or "", "TPNF ?")
    class_op = "%s | %s | %s" % (op, fin, tp)

    itens = []
    for det in dets:
        n_item = det.attrib.get("nItem", "")
        r = OrderedDict((c, None) for c in ITEM_COLS)
        r.update(meta)
        r.update(cab)
        r["NUMERO_ITEM"] = n_item
        if not n_item:
            log.add("AVISO", "EXTRACAO_ITEM", meta["ARQ_CAMINHO_COMPLETO"], "det sem atributo nItem",
                    "XML fora do padrao", "Item sem numero", "Item mantido com NUMERO_ITEM vazio")
        try:
            r["ID_REGISTRO"] = "%s-%03d" % (chave, int(n_item))
        except ValueError:
            r["ID_REGISTRO"] = "%s-%s" % (chave, n_item)
        prod = filho(det, "prod")
        imp = filho(det, "imposto")

        # Longa (todas as folhas do det)
        for cam, tag, val in folhas(det):
            sink_longa(r["ID_REGISTRO"], cam, tag, val)

        g = lambda t: txt(prod, t)
        r["COD_PRODUTO_FORNECEDOR"] = g("cProd")
        r["EAN"] = g("cEAN")
        r["EAN_TRIB"] = g("cEANTrib")
        r["DESCRICAO_PRODUTO"] = g("xProd")
        r["NCM"] = g("NCM")
        r["CEST"] = g("CEST")
        r["IND_ESCALA"] = g("indEscala")
        r["CNPJ_FABRICANTE"] = g("CNPJFab")
        r["COD_BENEF_FISCAL"] = g("cBenef")
        r["EXTIPI"] = g("EXTIPI")
        r["CFOP"] = g("CFOP")
        r["UNID_COMERCIAL"] = g("uCom")
        r["QTD_COMERCIAL"] = num(g("qCom"))
        r["VALOR_UNIT_COMERCIAL"] = num(g("vUnCom"))
        r["VALOR_PRODUTO"] = num(g("vProd"))
        r["UNID_TRIBUTAVEL"] = g("uTrib")
        r["QTD_TRIBUTAVEL"] = num(g("qTrib"))
        r["VALOR_UNIT_TRIBUTAVEL"] = num(g("vUnTrib"))
        r["VALOR_FRETE"] = num(g("vFrete"))
        r["VALOR_SEGURO"] = num(g("vSeg"))
        r["VALOR_DESCONTO"] = num(g("vDesc"))
        r["VALOR_OUTRAS_DESP"] = num(g("vOutro"))
        r["IND_TOTAL"] = g("indTot")
        r["PEDIDO_COMPRA"] = g("xPed")
        r["ITEM_PEDIDO"] = g("nItemPed")
        r["NUMERO_FCI"] = g("nFCI")
        med = filho(prod, "med")
        r["ANVISA_CODIGO"] = txt(med, "cProdANVISA") if med is not None else None
        r["ANVISA_MOTIVO_ISENCAO"] = txt(med, "xMotivoIsencao") if med is not None else None
        r["PMC_XML"] = num(txt(med, "vPMC")) if med is not None else None
        rastros = filhos(prod, "rastro")
        r["RASTRO_QTD_LOTES"] = len(rastros)
        r["RASTRO_LOTES"] = " / ".join("|".join(txt(rs, t) or "" for t in ("nLote", "qLote", "dFab", "dVal")) for rs in rastros) or None
        extras = [(local(c.tag), (c.text or "").strip()) for c in (prod if prod is not None else []) if local(c.tag) not in PROD_MAP]
        r["PROD_TAGS_NAO_MAPEADAS"] = pares_txt(extras) or None
        r["VALOR_TOTAL_ITEM_XML"] = num(txt(det, "vItem"))
        r["INF_AD_PROD"] = txt(det, "infAdProd")
        r["VALOR_TOT_TRIB_ITEM"] = num(txt(imp, "vTotTrib"))

        # ICMS
        icms_g = None
        icms_wrap = filho(imp, "ICMS")
        if icms_wrap is not None and len(icms_wrap):
            icms_g = list(icms_wrap)[0]
            if len(icms_wrap) > 1:
                log.add("AVISO", "EXTRACAO_ICMS", meta["ARQ_CAMINHO_COMPLETO"],
                        "Item %s com mais de um subgrupo ICMS" % n_item, "XML fora do padrao", "Somente o 1o mapeado em colunas",
                        "Demais preservados em 02B_TAGS_ITEM_LONGA")
        r["ICMS_GRUPO_TAG"] = local(icms_g.tag) if icms_g is not None else None
        icms_vals = {}
        if icms_g is not None:
            nao_map = []
            for c in icms_g:
                t = local(c.tag)
                icms_vals[t] = (c.text or "").strip()
                if t not in ICMS_MAP:
                    nao_map.append((t, icms_vals[t]))
            r["ICMS_TAGS_NAO_MAPEADAS"] = pares_txt(nao_map) or None
            r["ICMS_ORIGEM"] = icms_vals.get("orig")
            r["ICMS_CST"] = icms_vals.get("CST")
            r["ICMS_CSOSN"] = icms_vals.get("CSOSN")
            for t, _, tp in _ICMS_CAMPOS:
                if t in icms_vals:
                    r["ICMS_" + t] = num(icms_vals[t]) if tp == "NUMERO" else icms_vals[t]
        vI = lambda t: num(icms_vals.get(t))

        # DIFAL
        uf = filho(imp, "ICMSUFDest")
        if uf is not None:
            for c in uf:
                k = "DIFAL_" + local(c.tag)
                if k in r:
                    r[k] = num((c.text or "").strip())

        # IPI
        ipi = filho(imp, "IPI")
        ipi_sub = None
        if ipi is not None:
            r["IPI_COD_ENQ"] = txt(ipi, "cEnq")
            r["IPI_CNPJ_PRODUTOR"] = txt(ipi, "CNPJProd")
            r["IPI_COD_SELO"] = txt(ipi, "cSelo")
            r["IPI_QTD_SELO"] = num(txt(ipi, "qSelo"))
            ipi_sub = filho(ipi, "IPITrib")
            if ipi_sub is None:
                ipi_sub = filho(ipi, "IPINT")
            if ipi_sub is not None:
                r["IPI_GRUPO_TAG"] = local(ipi_sub.tag)
                r["IPI_CST"] = txt(ipi_sub, "CST")
                for t in ("vBC", "pIPI", "qUnid", "vUnid", "vIPI"):
                    r["IPI_" + t] = num(txt(ipi_sub, t))
        dev = filho(det, "impostoDevol")
        if dev is not None:
            r["IPI_DEVOL_PERC"] = num(txt(dev, "pDevol"))
            r["IPI_DEVOL_VALOR"] = num(txt(dev, "IPI", "vIPIDevol"))

        ii = filho(imp, "II")
        if ii is not None:
            for t in ("vBC", "vDespAdu", "vII", "vIOF"):
                r["II_" + t] = num(txt(ii, t))

        # PIS / COFINS
        for trib, mapa in (("PIS", PIS_MAP), ("COFINS", COFINS_MAP)):
            w = filho(imp, trib)
            if w is not None and len(w):
                sub = list(w)[0]
                r[trib + "_GRUPO_TAG"] = local(sub.tag)
                vals = {local(c.tag): (c.text or "").strip() for c in sub}
                r[trib + "_CST"] = vals.get("CST")
                r[trib + "_VBC"] = num(vals.get("vBC"))
                r[trib + "_ALIQ_PERC"] = num(vals.get("p" + trib))
                r[trib + "_QTD_BC_PROD"] = num(vals.get("qBCProd"))
                r[trib + "_ALIQ_VALOR"] = num(vals.get("vAliqProd"))
                r[trib + "_VALOR"] = num(vals.get("v" + trib))
                r[trib + "_TAGS_NAO_MAPEADAS"] = pares_txt([(k, v) for k, v in vals.items() if k not in mapa]) or None
            st = filho(imp, trib + "ST")
            if st is not None:
                r[trib + "ST_VALOR"] = num(txt(st, "v" + trib))

        # IBS/CBS
        ib = filho(imp, "IBSCBS")
        r["IBSCBS_PRESENTE"] = fmt_flag(ib is not None)
        if ib is not None:
            r["IBSCBS_CST"] = txt(ib, "CST")
            r["IBSCBS_CCLASSTRIB"] = txt(ib, "cClassTrib")
            g2 = filho(ib, "gIBSCBS")
            r["IBSCBS_VBC"] = num(txt(g2, "vBC"))
            r["IBS_UF_ALIQ"] = num(txt(g2, "gIBSUF", "pIBSUF"))
            r["IBS_UF_RED_ALIQ"] = num(txt(g2, "gIBSUF", "gRed", "pRedAliq"))
            r["IBS_UF_ALIQ_EFET"] = num(txt(g2, "gIBSUF", "gRed", "pAliqEfet"))
            r["IBS_UF_VALOR"] = num(txt(g2, "gIBSUF", "vIBSUF"))
            r["IBS_MUN_ALIQ"] = num(txt(g2, "gIBSMun", "pIBSMun"))
            r["IBS_MUN_RED_ALIQ"] = num(txt(g2, "gIBSMun", "gRed", "pRedAliq"))
            r["IBS_MUN_ALIQ_EFET"] = num(txt(g2, "gIBSMun", "gRed", "pAliqEfet"))
            r["IBS_MUN_VALOR"] = num(txt(g2, "gIBSMun", "vIBSMun"))
            r["IBS_VALOR"] = num(txt(g2, "vIBS"))
            r["CBS_ALIQ"] = num(txt(g2, "gCBS", "pCBS"))
            r["CBS_RED_ALIQ"] = num(txt(g2, "gCBS", "gRed", "pRedAliq"))
            r["CBS_ALIQ_EFET"] = num(txt(g2, "gCBS", "gRed", "pAliqEfet"))
            r["CBS_VALOR"] = num(txt(g2, "gCBS", "vCBS"))
            mapeadas = {"IBSCBS/CST", "IBSCBS/cClassTrib", "IBSCBS/gIBSCBS/vBC", "IBSCBS/gIBSCBS/gIBSUF/pIBSUF",
                        "IBSCBS/gIBSCBS/gIBSUF/gRed/pRedAliq", "IBSCBS/gIBSCBS/gIBSUF/gRed/pAliqEfet",
                        "IBSCBS/gIBSCBS/gIBSUF/vIBSUF", "IBSCBS/gIBSCBS/gIBSMun/pIBSMun", "IBSCBS/gIBSCBS/gIBSMun/vIBSMun",
                        "IBSCBS/gIBSCBS/gIBSMun/gRed/pRedAliq", "IBSCBS/gIBSCBS/gIBSMun/gRed/pAliqEfet",
                        "IBSCBS/gIBSCBS/vIBS", "IBSCBS/gIBSCBS/gCBS/pCBS", "IBSCBS/gIBSCBS/gCBS/gRed/pRedAliq",
                        "IBSCBS/gIBSCBS/gCBS/gRed/pAliqEfet", "IBSCBS/gIBSCBS/gCBS/vCBS"}
            resto = [(cam, v) for cam, _, v in folhas(ib) if cam not in mapeadas]
            r["IBSCBS_TAGS_NAO_MAPEADAS"] = pares_txt(resto) or None
        is_g = filho(imp, "IS")
        r["IS_PRESENTE"] = fmt_flag(is_g is not None)
        if is_g is not None:
            r["IS_TAGS"] = pares_txt([(c, v) for c, _, v in folhas(is_g)])
        conhecidos = {"ICMS", "IPI", "II", "PIS", "PISST", "COFINS", "COFINSST", "ICMSUFDest", "IBSCBS", "IS", "vTotTrib"}
        outros = [(local(c.tag), pares_txt([(cc, v) for cc, _, v in folhas(c)])) for c in (imp if imp is not None else [])
                  if local(c.tag) not in conhecidos]
        r["OUTROS_GRUPOS_IMPOSTO"] = " || ".join("%s: %s" % o for o in outros) or None

        # ---- Calculados
        if r["VALOR_PRODUTO"] is not None and r["QTD_COMERCIAL"] is not None and r["VALOR_UNIT_COMERCIAL"] is not None:
            r["CALC_DIF_VPROD_QCOM_X_VUNCOM"] = round(r["VALOR_PRODUTO"] - r["QTD_COMERCIAL"] * r["VALOR_UNIT_COMERCIAL"], 4)
        if vI("vICMS") is not None and vI("vBC") is not None and vI("pICMS") is not None:
            r["CALC_DIF_VICMS_BC_X_ALIQ"] = round(vI("vICMS") - vI("vBC") * vI("pICMS") / 100.0, 4)
        ean_ok = gtin_valido(r["EAN"] or "")
        r["EAN_GTIN_VALIDO"] = fmt_flag(ean_ok)
        if ean_ok:
            r["EAN_CHAVE_VINCULO"] = r["EAN"]
        elif gtin_valido(r["EAN_TRIB"] or ""):
            r["EAN_CHAVE_VINCULO"] = r["EAN_TRIB"]
        else:
            r["EAN_CHAVE_VINCULO"] = None

        # ---- Classificacao documental
        r["CLASS_DOC_TIPO_OPERACAO"] = class_op
        r["CLASS_DOC_ICMS"] = descreve_icms(r["ICMS_GRUPO_TAG"], r["ICMS_CST"] or "", r["ICMS_CSOSN"] or "", vI)
        r["FLAG_ST_DESTACADA"] = fmt_flag(gt0(vI("vICMSST")))
        r["FLAG_BC_ST_INFORMADA"] = fmt_flag(gt0(vI("vBCST")))
        r["FLAG_ST_RETIDA_ANTERIOR_INFO"] = fmt_flag(gt0(vI("vBCSTRet")) or gt0(vI("vICMSSTRet")))
        r["FLAG_ICMS_DESONERADO"] = fmt_flag(gt0(vI("vICMSDeson")))
        r["FLAG_DIFERIMENTO"] = fmt_flag(gt0(vI("pDif")) or gt0(vI("vICMSDif")))
        r["FLAG_FCP"] = fmt_flag(gt0(vI("vFCP")) or gt0(vI("vFCPST")) or gt0(vI("vFCPSTRet")))
        r["FLAG_IPI_VALOR"] = fmt_flag(gt0(r["IPI_vIPI"]))
        r["FLAG_DIFAL"] = fmt_flag(uf is not None)
        r["CLASS_DOC_PIS"] = descreve_piscofins("PIS", r["PIS_GRUPO_TAG"], r["PIS_CST"], r["PIS_VALOR"], r["PIS_ALIQ_PERC"])
        r["CLASS_DOC_COFINS"] = descreve_piscofins("COFINS", r["COFINS_GRUPO_TAG"], r["COFINS_CST"], r["COFINS_VALOR"], r["COFINS_ALIQ_PERC"])
        r["CLASS_DOC_IPI"] = descreve_ipi(r["IPI_GRUPO_TAG"], r["IPI_CST"], r["IPI_vIPI"])
        st_tag = "ST_DESTACADA" if r["FLAG_ST_DESTACADA"] == "S" else ("ST_RETIDA_INFO" if r["FLAG_ST_RETIDA_ANTERIOR_INFO"] == "S" else ("BC_ST_SEM_VALOR" if r["FLAG_BC_ST_INFORMADA"] == "S" else "SEM_ST_VALOR"))
        r["CENARIO_FISCAL_CHAVE"] = "|".join([
            "ICMS:%s/%s" % (r["ICMS_GRUPO_TAG"] or "-", r["ICMS_CST"] or r["ICMS_CSOSN"] or "-"),
            st_tag,
            "PIS:%s/%s" % (r["PIS_GRUPO_TAG"] or "-", r["PIS_CST"] or "-"),
            "COFINS:%s/%s" % (r["COFINS_GRUPO_TAG"] or "-", r["COFINS_CST"] or "-"),
            "IPI:%s/%s%s" % (r["IPI_GRUPO_TAG"] or "-", r["IPI_CST"] or "-", "+V" if r["FLAG_IPI_VALOR"] == "S" else ""),
            "IBSCBS:%s" % (r["IBSCBS_CST"] or "-"),
        ])

        # ---- Alertas documentais (checagens objetivas)
        al = []
        if r["FLAG_BC_ST_INFORMADA"] == "S" and r["FLAG_ST_DESTACADA"] == "N":
            al.append("BC_ST_INFORMADA_COM_VICMSST_ZERO")
        if r["ICMS_CST"] == "60" and r["FLAG_ST_RETIDA_ANTERIOR_INFO"] == "N":
            al.append("CST60_SEM_VALORES_ST_RETIDO")
        if r["PIS_CST"] in ("01", "02") and not gt0(r["PIS_VALOR"]):
            al.append("PIS_CST_TRIBUTADO_SEM_VALOR")
        if r["COFINS_CST"] in ("01", "02") and not gt0(r["COFINS_VALOR"]):
            al.append("COFINS_CST_TRIBUTADO_SEM_VALOR")
        if (r["PIS_CST"] or "") != (r["COFINS_CST"] or ""):
            al.append("CST_PIS_DIFERENTE_CST_COFINS")
        if r["CALC_DIF_VPROD_QCOM_X_VUNCOM"] is not None and abs(r["CALC_DIF_VPROD_QCOM_X_VUNCOM"]) > 0.011:
            al.append("VPROD_DIFERE_QCOM_X_VUNCOM")
        if r["CALC_DIF_VICMS_BC_X_ALIQ"] is not None and abs(r["CALC_DIF_VICMS_BC_X_ALIQ"]) > 0.02:
            al.append("VICMS_DIFERE_BC_X_ALIQ")
        ean_raw = r["EAN"] or ""
        if ean_raw in ("", "SEM GTIN"):
            al.append("EAN_AUSENTE_OU_SEM_GTIN")
        elif not ean_ok:
            al.append("EAN_INVALIDO_DV_OU_TAMANHO")
        if (r["EAN"] or "") != (r["EAN_TRIB"] or ""):
            al.append("EAN_DIFERENTE_EAN_TRIB")
        if (r["UNID_COMERCIAL"] or "") != (r["UNID_TRIBUTAVEL"] or ""):
            al.append("UNID_COMERCIAL_DIFERENTE_TRIBUTAVEL")
        if not re.fullmatch(r"\d{8}", r["NCM"] or ""):
            al.append("NCM_AUSENTE_OU_FORA_PADRAO")
        cf = r["CFOP"] or ""
        if cf[:1] in ("5", "1") and cab["ID_DEST"] == "2":
            al.append("CFOP_INTERNO_COM_IDDEST_INTERESTADUAL")
        if cf[:1] in ("6", "2") and cab["ID_DEST"] == "1":
            al.append("CFOP_INTERESTADUAL_COM_IDDEST_INTERNO")
        if cab["PROT_CSTAT"] is None:
            al.append("NOTA_SEM_PROTOCOLO_NO_ARQUIVO")
        elif cab["PROT_CSTAT"] not in ("100", "150"):
            al.append("PROTOCOLO_CSTAT_%s" % cab["PROT_CSTAT"])
        if cab["EVENTO_CANCELAMENTO_ENCONTRADO"] == "S":
            al.append("NOTA_COM_EVENTO_CANCELAMENTO")
        if r["ICMS_GRUPO_TAG"] is None:
            al.append("ITEM_SEM_GRUPO_ICMS")
        if r["PIS_GRUPO_TAG"] is None or r["COFINS_GRUPO_TAG"] is None:
            al.append("ITEM_SEM_GRUPO_PIS_OU_COFINS")
        r["ALERTAS_DOCUMENTAIS"] = " | ".join(al) or None

        # ---- Indicios (hipoteses, nao fatos)
        ind = []
        if r["PIS_CST"] == "04" or r["COFINS_CST"] == "04":
            ind.append("POSSIVEL MONOFASICO PIS/COFINS - REQUER VALIDACAO FISCAL")
        if r["PIS_CST"] == "06" or r["COFINS_CST"] == "06":
            ind.append("TRATAMENTO COMPATIVEL COM ALIQUOTA ZERO NO XML - REQUER VALIDACAO LEGAL")
        if r["PIS_CST"] == "05" or r["COFINS_CST"] == "05":
            ind.append("PIS/COFINS ST NO XML - REQUER VALIDACAO FISCAL")
        if r["PIS_CST"] in ("49", "99", "07", "08", "09") or r["COFINS_CST"] in ("49", "99", "07", "08", "09"):
            ind.append("PIS/COFINS SEM DEBITO NO XML (CST %s/%s) - ENQUADRAMENTO DO PRODUTO NAO DEMONSTRADO" % (r["PIS_CST"], r["COFINS_CST"]))
        if r["FLAG_ST_DESTACADA"] == "S" or r["ICMS_CST"] == "60" or r["FLAG_BC_ST_INFORMADA"] == "S":
            ind.append("PRODUTO COM INDICIO DE ICMS-ST NO DOCUMENTO - REGRA DE ST DA UF REQUER VALIDACAO")
        if "MESMA RAIZ" in class_op:
            ind.append("DOCUMENTO DE TRANSFERENCIA INTERNA - TRIBUTACAO PODE NAO REPRESENTAR A DA COMPRA")
        r["INDICIO_PARA_VALIDACAO"] = " | ".join(ind) or None
        r["STATUS_VALIDACAO_FISCAL"] = "PENDENTE"
        for k in INTERNAR:
            if isinstance(r[k], str):
                r[k] = sys.intern(r[k])
        itens.append(Row(r))
    return itens


# ----------------------------------------------------------------------------
# Varredura de arquivos
# ----------------------------------------------------------------------------

def iterar_arquivos(origem, excluir_nomes, saida_abs):
    excl = {e.lower() for e in excluir_nomes}
    for dirpath, dirnames, filenames in os.walk(origem):
        dirnames[:] = sorted(d for d in dirnames
                             if d.lower() not in excl and os.path.abspath(os.path.join(dirpath, d)) != saida_abs)
        for f in sorted(filenames):
            yield os.path.join(dirpath, f)


def reler_bytes(caminho_logico):
    """Rele (somente leitura) um XML solto ou membro de ZIP/RAR pelo caminho logico 'arquivo::membro'."""
    if "::" not in caminho_logico:
        with open(caminho_logico, "rb") as fh:
            return fh.read()
    cont, membro = caminho_logico.split("::", 1)
    if cont.lower().endswith(".zip"):
        with zipfile.ZipFile(cont, "r") as z:
            return z.read(membro)
    import rarfile
    with rarfile.RarFile(cont) as rz:
        return rz.read(membro)


def conteudos(caminho_arq, ler_zip, log):
    """Gera (caminho_logico, nome, extensao, tamanho, mtime, bytes|None, tipo_container, erro)."""
    ext = os.path.splitext(caminho_arq)[1].lower()
    st = os.stat(caminho_arq)
    mtime = dt.datetime.fromtimestamp(st.st_mtime).strftime("%Y-%m-%d %H:%M:%S")
    nome = os.path.basename(caminho_arq)
    if ext == ".xml":
        with open(caminho_arq, "rb") as fh:  # somente leitura
            yield caminho_arq, nome, ext, st.st_size, mtime, fh.read(), "ARQUIVO", None
        return
    if ext == ".zip" and ler_zip:
        try:
            with zipfile.ZipFile(caminho_arq, "r") as z:
                yield caminho_arq, nome, ext, st.st_size, mtime, None, "COMPACTADO_PROCESSADO", None
                for info in z.infolist():
                    if info.filename.lower().endswith(".xml") and not info.is_dir():
                        data = z.read(info)
                        yield (caminho_arq + "::" + info.filename, os.path.basename(info.filename), ".xml",
                               info.file_size, "%04d-%02d-%02d %02d:%02d:%02d" % info.date_time, data, "DENTRO_ZIP", None)
        except Exception as e:
            yield caminho_arq, nome, ext, st.st_size, mtime, None, "COMPACTADO_NAO_PROCESSADO", "ZIP ilegivel: %s" % e
        return
    if ext == ".rar" and ler_zip:
        try:
            import rarfile  # opcional
            with rarfile.RarFile(caminho_arq) as rz:
                yield caminho_arq, nome, ext, st.st_size, mtime, None, "COMPACTADO_PROCESSADO", None
                for info in rz.infolist():
                    if info.filename.lower().endswith(".xml") and not info.is_dir():
                        data = rz.read(info)
                        yield (caminho_arq + "::" + info.filename, os.path.basename(info.filename), ".xml",
                               info.file_size, "", data, "DENTRO_RAR", None)
        except ImportError:
            yield caminho_arq, nome, ext, st.st_size, mtime, None, "COMPACTADO_NAO_PROCESSADO", \
                "Biblioteca rarfile/unrar indisponivel - conteudo do RAR nao lido"
        except Exception as e:
            yield caminho_arq, nome, ext, st.st_size, mtime, None, "COMPACTADO_NAO_PROCESSADO", "RAR ilegivel: %s" % e
        return
    yield caminho_arq, nome, ext, st.st_size, mtime, None, "IGNORADO_NAO_XML", None


# ----------------------------------------------------------------------------
# Divergencias, cenarios, amostra
# ----------------------------------------------------------------------------

CAMPOS_DIVERG = ["NCM", "CEST", "CFOP", "ICMS_GRUPO_TAG", "ICMS_CST", "ICMS_CSOSN", "ICMS_ORIGEM", "FLAG_ST_DESTACADA",
                 "FLAG_ST_RETIDA_ANTERIOR_INFO", "ICMS_pICMS", "ICMS_pRedBC", "ICMS_pMVAST", "ICMS_pRedBCST", "ICMS_pICMSST",
                 "PIS_GRUPO_TAG", "PIS_CST", "PIS_ALIQ_PERC", "COFINS_GRUPO_TAG", "COFINS_CST", "COFINS_ALIQ_PERC",
                 "IPI_GRUPO_TAG", "IPI_CST", "IPI_pIPI", "IBSCBS_CST", "IBSCBS_CCLASSTRIB", "COD_BENEF_FISCAL",
                 "CENARIO_FISCAL_CHAVE"]


def vstr(v):
    if v is None:
        return "(AUSENTE)"
    if isinstance(v, float):
        return ("%.4f" % v).rstrip("0").rstrip(".")
    return str(v) if v != "" else "(VAZIO)"


def calc_divergencias(itens):
    por_ean = defaultdict(list)
    for r in itens:
        if r["EAN_CHAVE_VINCULO"]:
            por_ean[r["EAN_CHAVE_VINCULO"]].append(r)
    linhas = []
    eans_div = set()
    for ean, regs in por_ean.items():
        if len(regs) < 2:
            continue
        for campo in CAMPOS_DIVERG:
            dist = defaultdict(list)
            for r in regs:
                dist[vstr(r[campo])].append(r)
            if len(dist) < 2:
                continue
            eans_div.add(ean)
            fornecedores = {r["EMIT_CNPJ"] for r in regs}
            ops = {("INTERNA" if "MESMA RAIZ" in (r["CLASS_DOC_TIPO_OPERACAO"] or "") else "TERCEIRO") for r in regs}
            # divergencia dentro do mesmo fornecedor?
            mesmo_forn = any(len({vstr(r[campo]) for r in regs if r["EMIT_CNPJ"] == f}) > 1 for f in fornecedores)
            so_terc = [r for r in regs if "MESMA RAIZ" not in (r["CLASS_DOC_TIPO_OPERACAO"] or "")]
            div_so_terc = len({vstr(r[campo]) for r in so_terc}) > 1
            linhas.append(OrderedDict([
                ("EAN_CHAVE_VINCULO", ean),
                ("DESCRICAO_EXEMPLO", regs[0]["DESCRICAO_PRODUTO"]),
                ("CAMPO", campo),
                ("QTD_VALORES_DISTINTOS", len(dist)),
                ("VALORES_E_QTD", " ; ".join("%s (%d)" % (k, len(v)) for k, v in sorted(dist.items()))),
                ("ID_REGISTROS_POR_VALOR", " ; ".join("%s: %s" % (k, ",".join(x["ID_REGISTRO"] for x in v[:5]) + ("..." if len(v) > 5 else ""))
                                                       for k, v in sorted(dist.items()))),
                ("QTD_REGISTROS_EAN", len(regs)),
                ("QTD_FORNECEDORES", len(fornecedores)),
                ("FORNECEDORES", " ; ".join(sorted({"%s %s" % (r["EMIT_CNPJ"], r["EMIT_RAZAO_SOCIAL"]) for r in regs}))),
                ("TIPOS_OPERACAO_ENVOLVIDOS", " + ".join(sorted(ops))),
                ("DIVERGE_DENTRO_DO_MESMO_FORNECEDOR", fmt_flag(mesmo_forn)),
                ("DIVERGE_SO_ENTRE_COMPRAS_DE_TERCEIROS", fmt_flag(div_so_terc)),
                ("STATUS", "REGISTRADO - NAO CORRIGIDO - REQUER ANALISE"),
            ]))
    for r in itens:
        r["EAN_COM_DIVERGENCIA"] = fmt_flag(r["EAN_CHAVE_VINCULO"] in eans_div)
    rank = {}
    for l in linhas:
        e = l["EAN_CHAVE_VINCULO"]
        mf, nc, nr = rank.get(e, (0, 0, 0))
        rank[e] = (max(mf, 1 if l["DIVERGE_DENTRO_DO_MESMO_FORNECEDOR"] == "S" else 0), nc + 1, l["QTD_REGISTROS_EAN"])
    return linhas, rank


def calc_cenarios(itens):
    grupos = OrderedDict()
    for r in sorted(itens, key=lambda x: x["CENARIO_FISCAL_CHAVE"]):
        grupos.setdefault(r["CENARIO_FISCAL_CHAVE"], []).append(r)
    ordem = sorted(grupos.items(), key=lambda kv: -len(kv[1]))
    mapa = []
    ids = {}
    for i, (k, regs) in enumerate(ordem, 1):
        cid = "CEN%03d" % i
        ids[k] = cid
        ex = regs[0]
        cnt = lambda campo: Counter(vstr(r[campo]) for r in regs)
        mapa.append(OrderedDict([
            ("CENARIO_FISCAL_ID", cid),
            ("CENARIO_FISCAL_CHAVE", k),
            ("CLASS_DOC_ICMS", " ; ".join(sorted({r["CLASS_DOC_ICMS"] for r in regs}))),
            ("CLASS_DOC_PIS", ex["CLASS_DOC_PIS"]),
            ("CLASS_DOC_COFINS", ex["CLASS_DOC_COFINS"]),
            ("CLASS_DOC_IPI", " ; ".join(sorted({r["CLASS_DOC_IPI"] for r in regs}))),
            ("QTD_ITENS", len(regs)),
            ("QTD_NOTAS", len({r["CHAVE_NFE"] for r in regs})),
            ("QTD_FORNECEDORES", len({r["EMIT_CNPJ"] for r in regs})),
            ("QTD_EAN", len({r["EAN_CHAVE_VINCULO"] for r in regs if r["EAN_CHAVE_VINCULO"]})),
            ("QTD_NCM", len({r["NCM"] for r in regs})),
            ("CFOPS", " ; ".join("%s (%d)" % kv for kv in cnt("CFOP").most_common())),
            ("ORIGENS", " ; ".join("%s (%d)" % kv for kv in cnt("ICMS_ORIGEM").most_common())),
            ("NCM_TOP10", " ; ".join("%s (%d)" % kv for kv in cnt("NCM").most_common(10))),
            ("TIPOS_OPERACAO", " ; ".join("%s (%d)" % kv for kv in cnt("CLASS_DOC_TIPO_OPERACAO").most_common())),
            ("INDICIOS_PARA_VALIDACAO", " ; ".join(sorted({r["INDICIO_PARA_VALIDACAO"] or "-" for r in regs}))),
            ("QTD_AMOSTRAS_SELECIONADAS", 0),
            ("STATUS_VALIDACAO_FISCAL", "PENDENTE"),
        ]))
    for r in itens:
        r["CENARIO_FISCAL_ID"] = ids[r["CENARIO_FISCAL_CHAVE"]]
    return mapa


def completude(r):
    pts = 0
    pts += 2 if r["EAN_GTIN_VALIDO"] == "S" else 0
    pts += 1 if r["CEST"] else 0
    pts += 1 if r["NCM"] else 0
    pts += 1 if r["PROT_CSTAT"] == "100" else 0
    pts += 1 if not r["ALERTAS_DOCUMENTAIS"] else 0
    pts += 1 if "TERCEIRO" in (r["CLASS_DOC_TIPO_OPERACAO"] or "") else 0
    return pts


class LinhaAmostra(object):
    """Linha da amostra por referencia ao item (sem copiar as 245 colunas)."""
    __slots__ = ("seq", "motivo", "row")

    def __init__(self, seq, motivo, row):
        self.seq, self.motivo, self.row = seq, motivo, row

    def get(self, k, default=None):
        if k == "AMOSTRA_SEQ":
            return self.seq
        if k == "MOTIVO_SELECAO":
            return self.motivo
        return self.row.get(k, default)

    def __getitem__(self, k):
        return self.get(k)


def selecionar_amostra(itens, mapa, rank_div, n_por_cenario, max_eans_div=150):
    sel = OrderedDict()  # ID_REGISTRO -> set(motivos)

    def add(r, motivo):
        sel.setdefault(r["ID_REGISTRO"], [r, []])[1].append(motivo)

    # 1) Diversidade por cenario (forn/CFOP/NCM novos primeiro, depois completude)
    por_cen = defaultdict(list)
    for r in itens:
        por_cen[r["CENARIO_FISCAL_ID"]].append(r)
    for cid, regs in por_cen.items():
        escolhidos = []
        vistos = {"f": set(), "c": set(), "n": set(), "e": set()}
        candidatos = sorted(regs, key=lambda r: (-completude(r), r["ID_REGISTRO"]))
        while len(escolhidos) < n_por_cenario and candidatos:
            def ganho(r):
                return ((r["EMIT_CNPJ"] not in vistos["f"]) * 4 + (r["CFOP"] not in vistos["c"]) * 2 +
                        (r["NCM"] not in vistos["n"]) * 1 + (r["EAN_CHAVE_VINCULO"] not in vistos["e"]) * 1,
                        completude(r))
            melhor = max(candidatos, key=ganho)
            if escolhidos and ganho(melhor)[0] == 0:
                break  # evita registros identicos quando nao ha diversidade
            candidatos.remove(melhor)
            escolhidos.append(melhor)
            vistos["f"].add(melhor["EMIT_CNPJ"]); vistos["c"].add(melhor["CFOP"])
            vistos["n"].add(melhor["NCM"]); vistos["e"].add(melhor["EAN_CHAVE_VINCULO"])
        for i, r in enumerate(escolhidos, 1):
            add(r, "CENARIO %s (exemplo %d/%d do cenario; %d itens no cenario)" % (cid, i, len(escolhidos), len(regs)))

    # 2) Cobertura de CFOP, origem, fornecedor, grupos ICMS/PIS/COFINS/IPI
    for campo, rotulo in (("CFOP", "CFOP"), ("ICMS_ORIGEM", "ORIGEM"), ("EMIT_CNPJ", "FORNECEDOR"),
                          ("ICMS_GRUPO_TAG", "GRUPO ICMS"), ("PIS_GRUPO_TAG", "GRUPO PIS"),
                          ("COFINS_GRUPO_TAG", "GRUPO COFINS"), ("IPI_GRUPO_TAG", "GRUPO IPI"),
                          ("IBSCBS_CST", "CST IBSCBS"), ("CEST", "CEST")):
        ja = {vstr(v[0][campo]) for v in sel.values()}
        por_val = defaultdict(list)
        for r in itens:
            por_val[vstr(r[campo])].append(r)
        for val, regs in por_val.items():
            if val not in ja:
                r = max(regs, key=completude)
                add(r, "COBERTURA %s=%s (%d itens)" % (rotulo, val, len(regs)))

    # 3) Produtos recorrentes com divergencia: prioriza divergencia dentro do mesmo fornecedor,
    #    mais campos divergentes e mais registros; ate max_eans_div EANs e 2 variantes por EAN
    #    (as variantes mais diferentes entre si). Lista completa permanece em 04_DIVERGENCIAS_PRODUTOS.
    priorizados = sorted(rank_div, key=lambda e: (-rank_div[e][0], -rank_div[e][1], -rank_div[e][2], e))[:max_eans_div]
    alvo = set(priorizados)
    idx_ean = defaultdict(list)
    for r in itens:
        if r["EAN_CHAVE_VINCULO"] in alvo:
            idx_ean[r["EAN_CHAVE_VINCULO"]].append(r)
    for pos, ean in enumerate(priorizados, 1):
        regs = idx_ean[ean]
        variantes = defaultdict(list)
        for r in regs:
            variantes[r["CENARIO_FISCAL_CHAVE"] + "|" + (r["NCM"] or "") + "|" + (r["CFOP"] or "")].append(r)
        reps_var = [max(rr, key=completude) for rr in variantes.values()]
        if len(reps_var) > 2:
            def dist(a, b):
                return sum(1 for c in CAMPOS_DIVERG if vstr(a[c]) != vstr(b[c]))
            par = max(((a, b) for i, a in enumerate(reps_var) for b in reps_var[i + 1:]), key=lambda ab: dist(*ab))
            reps_var = list(par)
        for r in reps_var:
            add(r, "DIVERGENCIA EAN %s (prioridade %d; %d campos divergentes; %d variantes)" % (ean, pos, rank_div[ean][1], len(variantes)))

    # 4) Atipicos: 1 exemplo por tipo de alerta (ate 2 fornecedores diferentes)
    por_alerta = defaultdict(list)
    for r in itens:
        for a in (r["ALERTAS_DOCUMENTAIS"] or "").split(" | "):
            if a:
                por_alerta[a].append(r)
    for a, regs in por_alerta.items():
        forn = set()
        for r in sorted(regs, key=lambda x: -completude(x)):
            if r["EMIT_CNPJ"] in forn:
                continue
            forn.add(r["EMIT_CNPJ"])
            add(r, "ATIPICO %s (%d itens com o alerta)" % (a, len(regs)))
            if len(forn) >= 2:
                break

    linhas = []
    cont_cen = Counter()
    for i, (rid, (r, motivos)) in enumerate(sel.items(), 1):
        linhas.append(LinhaAmostra(i, " | ".join(motivos), r))
        cont_cen[r["CENARIO_FISCAL_ID"]] += 1
    for m in mapa:
        m["QTD_AMOSTRAS_SELECIONADAS"] = cont_cen.get(m["CENARIO_FISCAL_ID"], 0)
    return linhas


# ----------------------------------------------------------------------------
# Escrita
# ----------------------------------------------------------------------------

def escrever_csv(caminho_csv, cabecalho, linhas):
    with open(caminho_csv, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh, delimiter=";", quoting=csv.QUOTE_MINIMAL)
        w.writerow(cabecalho)
        for l in linhas:
            w.writerow(["" if l.get(c) is None else (("%s" % l.get(c)).replace(".", ",") if isinstance(l.get(c), float) else l.get(c))
                        for c in cabecalho])


def escrever_xlsx(caminho_xlsx, abas, log, tipos):
    try:
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        log.add("AVISO", "SAIDA_XLSX", caminho_xlsx, "openpyxl nao instalado", "Biblioteca ausente",
                "XLSX nao gerado", "Somente CSVs gerados (pip install openpyxl)")
        return False
    wb = Workbook(write_only=True)
    bold = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F4E78")
    for nome, cab, linhas in abas:
        ws = wb.create_sheet(title=nome[:31])
        ws.freeze_panes = "A2"
        for i, c in enumerate(cab, 1):
            ws.column_dimensions[get_column_letter(i)].width = min(max(12, len(c) + 2), 45)
        hdr = []
        for c in cab:
            cell = WriteOnlyCell(ws, value=c)
            cell.font = bold
            cell.fill = fill
            hdr.append(cell)
        ws.append(hdr)
        if len(linhas) > XLSX_MAX_LINHAS:
            log.add("AVISO", "SAIDA_XLSX", nome, "Aba excede limite do Excel (%d linhas)" % len(linhas),
                    "Volume", "Aba truncada no XLSX", "Usar o CSV completo correspondente")
            linhas = linhas[:XLSX_MAX_LINHAS]
        for l in linhas:
            row = []
            for c in cab:
                v = l.get(c)
                if v is None:
                    row.append(None)
                    continue
                if tipos.get(c) in ("TEXTO", "DATA_TEXTO", "FLAG") or not isinstance(v, (int, float)):
                    # gravado como celula de texto (tipo 's'): Excel nao converte nem remove zeros a esquerda
                    v = str(v)
                    if len(v) > XLSX_MAX_CEL:
                        v = v[:XLSX_MAX_CEL] + " [TRUNCADO NO XLSX - VER CSV]"
                    if v[:1] == "=":
                        cell = WriteOnlyCell(ws, value=v)
                        cell.data_type = "s"  # texto literal, nunca formula
                        row.append(cell)
                    else:
                        row.append(v)
                else:
                    row.append(v)
            ws.append(row)
    wb.save(caminho_xlsx)
    return True


# ----------------------------------------------------------------------------
# Principal
# ----------------------------------------------------------------------------

def main(argv=None):
    ap = argparse.ArgumentParser(description="Extrator base fiscal de XML NF-e (somente leitura).")
    ap.add_argument("--origem", required=True)
    ap.add_argument("--saida", required=True)
    ap.add_argument("--excluir", action="append", default=[], help="Nome de pasta a ignorar (repetivel)")
    ap.add_argument("--amostra-por-cenario", type=int, default=3)
    ap.add_argument("--max-eans-divergentes-amostra", type=int, default=150)
    ap.add_argument("--sem-zip", action="store_true", help="Nao ler .zip/.rar")
    ap.add_argument("--sem-xlsx", action="store_true")
    a = ap.parse_args(argv)

    inicio = dt.datetime.now()
    data_proc = inicio.strftime("%Y-%m-%d %H:%M:%S")
    origem = os.path.abspath(a.origem)
    if not os.path.isdir(origem):
        print("ERRO: pasta de origem inacessivel: %s" % origem)
        return 2
    exec_dir = os.path.join(os.path.abspath(a.saida), "EXEC_" + inicio.strftime("%Y%m%d_%H%M%S"))
    if os.path.commonpath([exec_dir, origem]) == origem:
        excluir_auto = os.path.abspath(a.saida)
    else:
        excluir_auto = ""
    os.makedirs(exec_dir, exist_ok=False)  # nunca sobrescreve execucao anterior
    log = Log()

    def progresso(msg):
        print("[%s] %s" % (dt.datetime.now().strftime("%H:%M:%S"), msg), flush=True)

    progresso("Inicio. Origem: %s" % origem)

    # --- Passo 1: indice leve de todos os arquivos (nenhuma arvore XML fica em memoria)
    inv = []          # registros de inventario
    info = {}         # id(rec) -> dict(tipo_raiz, tem_prot)
    eventos_cancel = set()

    def registrar(cam, nome, ext, tam, mtime, data, cont, erro):
        rec = OrderedDict((c, None) for c, _ in INV_COLS)
        rec["CAMINHO_COMPLETO"] = cam
        rec["CAMINHO_RELATIVO"] = os.path.relpath(cam.split("::")[0], origem) + (("::" + cam.split("::", 1)[1]) if "::" in cam else "")
        rec["NOME_ARQUIVO"] = nome
        rec["EXTENSAO"] = ext
        rec["TAMANHO_BYTES"] = tam
        rec["DATA_MODIFICACAO"] = mtime
        rec["DATA_PROCESSAMENTO"] = data_proc
        rec["QTD_ITENS_EXTRAIDOS"] = 0
        inf_r = {"raiz": None, "prot": False}
        if data is None:
            rec["TIPO_DOCUMENTO"] = "COMPACTADO" if cont.startswith("COMPACTADO") else ("ERRO_ACESSO" if cont == "ERRO_ACESSO" else "NAO_XML")
            rec["STATUS"] = cont if cont != "ARQUIVO" else "IGNORADO_NAO_XML"
            rec["ERRO"] = erro
            if cont == "COMPACTADO_NAO_PROCESSADO":
                log.add("AVISO", "INVENTARIO", cam, erro or "Compactado nao processado", "Arquivo compactado",
                        "XMLs internos nao entram na base nesta execucao",
                        "Registrado; extrair para pasta propria ou instalar rarfile/unrar e reexecutar")
        else:
            rec["HASH_SHA256"] = sha256_bytes(data)
            try:
                root, aviso = parse_xml_bytes(data)
                rec["TIPO_DOCUMENTO"] = classificar_documento(root)
                inf_r["raiz"] = local(root.tag)
                inf_r["prot"] = caminho(root, "protNFe", "infProt") is not None
                if aviso:
                    rec["ERRO"] = aviso
                    log.add("AVISO", "LEITURA_XML", cam, aviso, "Encoding/BOM fora do padrao", "Lido em memoria com ajuste",
                            "Original preservado sem alteracao")
                if rec["TIPO_DOCUMENTO"] == "procEventoNFe":
                    inf_ev = next((el for el in root.iter() if local(el.tag) == "infEvento"), None)
                    tp = txt(inf_ev, "tpEvento") if inf_ev is not None else None
                    ch = txt(inf_ev, "chNFe") if inf_ev is not None else None
                    rec["EVENTO_TIPO"] = tp
                    rec["CHAVE_NFE"] = ch
                    rec["STATUS"] = "EVENTO"
                    if tp == "110111" and ch:
                        eventos_cancel.add(ch)
                elif rec["TIPO_DOCUMENTO"] in ("nfeProc", "NFe_SEM_PROTOCOLO"):
                    nfe = root if local(root.tag) == "NFe" else filho(root, "NFe")
                    inf = filho(nfe, "infNFe")
                    if inf is None:
                        raise ValueError("infNFe nao encontrado")
                    ide = filho(inf, "ide")
                    ch = (inf.attrib.get("Id") or "").replace("NFe", "")
                    rec["CHAVE_NFE"] = ch
                    rec["MODELO"] = txt(ide, "mod")
                    rec["NUMERO_NF"] = txt(ide, "nNF")
                    rec["SERIE"] = txt(ide, "serie")
                    rec["DATA_EMISSAO"] = txt(ide, "dhEmi") or txt(ide, "dEmi")
                    rec["EMIT_CNPJ"] = txt(inf, "emit", "CNPJ")
                    rec["EMIT_RAZAO_SOCIAL"] = txt(inf, "emit", "xNome")
                    rec["DEST_CNPJ"] = txt(inf, "dest", "CNPJ")
                    rec["DEST_RAZAO_SOCIAL"] = txt(inf, "dest", "xNome")
                    rec["NATUREZA_OPERACAO"] = txt(ide, "natOp")
                    rec["QTD_ITENS"] = len(filhos(inf, "det"))
                    rec["PROT_CSTAT"] = txt(root, "protNFe", "infProt", "cStat")
                    rec["CHAVE_DV_OK"] = fmt_flag(chave_dv_ok(ch))
                    m = re.findall(r"\d{44}", rec["NOME_ARQUIVO"] or "")
                    rec["CHAVE_IGUAL_NOME_ARQUIVO"] = (fmt_flag(ch in m) if m else None)
                    if rec["CHAVE_DV_OK"] == "N":
                        log.add("AVISO", "VALIDACAO_CHAVE", cam, "DV da chave nao confere: %s" % ch,
                                "Chave invalida ou XML alterado", "Rastreabilidade comprometida", "Mantido; sinalizado")
                    if rec["CHAVE_IGUAL_NOME_ARQUIVO"] == "N":
                        log.add("AVISO", "VALIDACAO_CHAVE", cam, "Nome do arquivo contem chave diferente da interna",
                                "Arquivo renomeado/trocado", "Possivel confusao de vinculo", "Mantido; usa chave interna")
                else:
                    rec["STATUS"] = "NAO_NFE"
            except Exception as e:
                rec["TIPO_DOCUMENTO"] = rec["TIPO_DOCUMENTO"] or "XML_ILEGIVEL"
                rec["STATUS"] = "ERRO_LEITURA"
                rec["ERRO"] = str(e)
                log.add("ERRO", "LEITURA_XML", cam, str(e), "XML corrompido, truncado, estrutura NF-e incompleta ou nao-XML com extensao .xml",
                        "Arquivo fora da base", "Registrado no inventario com STATUS=ERRO_LEITURA")
        inv.append(rec)
        info[id(rec)] = inf_r

    for p in iterar_arquivos(origem, a.excluir, excluir_auto):
        try:
            for item in conteudos(p, not a.sem_zip, log):
                registrar(*item)
        except Exception as e:
            registrar(p, os.path.basename(p), os.path.splitext(p)[1].lower(), None, None, None, "ERRO_ACESSO", str(e))
            log.add("ERRO", "INVENTARIO", p, str(e), "Permissao/arquivo em uso/sincronizacao", "Arquivo nao lido", "Registrado no inventario")

    progresso("Passo 1 concluido: %d arquivos inventariados" % len(inv))
    # --- Passo 2: canonico por hash e por chave (ordem deterministica)
    def prioridade(rec):
        ir = info[id(rec)]
        return (0 if ir["raiz"] == "nfeProc" else 1, 0 if ir["prot"] else 1,
                1 if "::" in rec["CAMINHO_COMPLETO"] else 0, len(rec["CAMINHO_COMPLETO"]), rec["CAMINHO_COMPLETO"])

    candidatos = defaultdict(list)
    hash_visto = {}
    for rec in sorted(inv, key=prioridade):
        if rec["STATUS"] is not None or rec["TIPO_DOCUMENTO"] not in ("nfeProc", "NFe_SEM_PROTOCOLO"):
            continue
        if rec["HASH_SHA256"] in hash_visto:
            rec["STATUS"] = "DUPLICADO_HASH"
            rec["DUPLICADO_DE"] = hash_visto[rec["HASH_SHA256"]]
            continue
        hash_visto[rec["HASH_SHA256"]] = rec["CAMINHO_COMPLETO"]
        candidatos[rec["CHAVE_NFE"]].append(rec)

    dup_hash_por_chave = Counter(r["CHAVE_NFE"] for r in inv if r["STATUS"] == "DUPLICADO_HASH")
    qtd_por_chave = {}
    canonicos = []
    for ch, lst in candidatos.items():
        qtd_por_chave[ch] = len(lst) + dup_hash_por_chave.get(ch, 0)
        canonicos.append(lst[0])
        for rec in lst[1:]:
            rec["STATUS"] = "DUPLICADO_CHAVE"
            rec["DUPLICADO_DE"] = lst[0]["CAMINHO_COMPLETO"]
    canonicos.sort(key=lambda r: r["CAMINHO_COMPLETO"])

    progresso("Passo 2 concluido: %d notas canonicas" % len(canonicos))
    # --- Passo 3: extracao por item (rele somente os canonicos; tabela longa gravada em streaming)
    itens = []
    obs = Counter()
    obs_ex = {}
    tot_long = 0
    fh_long = open(os.path.join(exec_dir, "02B_TAGS_ITEM_LONGA.csv"), "w", newline="", encoding="utf-8-sig")
    w_long = csv.writer(fh_long, delimiter=";")
    w_long.writerow(["ID_REGISTRO", "CAMINHO_TAG", "TAG", "VALOR_ORIGINAL"])

    def sink_longa(rid, cam, tag, val):
        nonlocal tot_long
        w_long.writerow([rid, cam, tag, val])
        tot_long += 1
        partes = cam.split("/")
        k = ("/".join(partes[1:-1]) or "(det)", partes[-1])  # ex.: imposto/IBSCBS/gIBSCBS/gIBSUF/gRed
        obs[k] += 1
        if k not in obs_ex:
            obs_ex[k] = rid

    for rec in canonicos:
        loja, tipo = info_pasta(rec["CAMINHO_RELATIVO"])
        meta = OrderedDict([
            ("ARQ_CAMINHO_COMPLETO", rec["CAMINHO_COMPLETO"]),
            ("ARQ_CAMINHO_RELATIVO", rec["CAMINHO_RELATIVO"]),
            ("ARQ_NOME", rec["NOME_ARQUIVO"]),
            ("ARQ_PASTA_LOJA_CNPJ", loja or None),
            ("ARQ_PASTA_TIPO", tipo or None),
            ("ARQ_HASH_SHA256", rec["HASH_SHA256"]),
            ("QTD_ARQUIVOS_MESMA_CHAVE", qtd_por_chave.get(rec["CHAVE_NFE"], 1)),
        ])
        try:
            data = reler_bytes(rec["CAMINHO_COMPLETO"])
            if sha256_bytes(data) != rec["HASH_SHA256"]:
                raise ValueError("Arquivo mudou entre a leitura do inventario e a extracao (hash diferente)")
            root, _ = parse_xml_bytes(data)
            del data
            its = extrair_nfe(root, meta, log, eventos_cancel, sink_longa)
            del root
            itens.extend(its)
            rec["QTD_ITENS_EXTRAIDOS"] = len(its)
            rec["STATUS"] = "LIDO_COM_AVISO" if rec["ERRO"] else "LIDO_OK"
            if len(its) != rec["QTD_ITENS"]:
                log.add("ERRO", "CONFERENCIA_ITENS", rec["CAMINHO_COMPLETO"], "Itens extraidos (%d) <> det (%d)" % (len(its), rec["QTD_ITENS"]),
                        "Falha de extracao", "Base incompleta para a nota", "Verificar manualmente")
        except Exception as e:
            rec["STATUS"] = "ERRO_LEITURA"
            rec["ERRO"] = "Falha na extracao: %s" % e
            log.add("ERRO", "EXTRACAO", rec["CAMINHO_COMPLETO"], str(e), "Estrutura inesperada ou arquivo alterado durante a execucao",
                    "Nota fora da base", "Registrado")
    fh_long.close()

    progresso("Passo 3 concluido: %d itens extraidos" % len(itens))
    # --- Passo 4: divergencias, cenarios, amostra
    divergencias, rank_div = calc_divergencias(itens)
    eans_div = set(rank_div)
    mapa = calc_cenarios(itens)
    amostra = selecionar_amostra(itens, mapa, rank_div, a.amostra_por_cenario, a.max_eans_divergentes_amostra)

    # Grupos/tags observados (inspecao de estrutura)
    grupos_obs = [OrderedDict([("GRUPO_CAMINHO", k[0]), ("TAG", k[1]), ("OCORRENCIAS", v), ("EXEMPLO_ID_REGISTRO", obs_ex[k])])
                  for k, v in sorted(obs.items())]

    # --- Passo 5: validacoes
    stat = Counter(r["STATUS"] for r in inv)
    chk = []

    def C(cat, nome, valor, ok="INFO", det=""):
        chk.append(OrderedDict([("CATEGORIA", cat), ("CHECAGEM", nome), ("RESULTADO", valor), ("STATUS", ok), ("DETALHE", det)]))

    xml_total = sum(1 for r in inv if r["EXTENSAO"] == ".xml")
    C("QUANTITATIVA", "Arquivos encontrados (todos, inclusive dentro de compactados lidos)", len(inv))
    C("QUANTITATIVA", "Arquivos .xml encontrados", xml_total)
    for s, n in sorted(stat.items(), key=lambda kv: str(kv[0])):
        C("QUANTITATIVA", "Arquivos com STATUS=%s" % s, n, "ATENCAO" if s in ("ERRO_LEITURA", "COMPACTADO_NAO_PROCESSADO") else "INFO")
    C("QUANTITATIVA", "Notas distintas na base (chaves canonicas)", len(canonicos))
    C("QUANTITATIVA", "Itens extraidos", len(itens))
    C("QUANTITATIVA", "Fornecedores (EMIT_CNPJ distintos)", len({r["EMIT_CNPJ"] for r in itens}))
    C("QUANTITATIVA", "Destinatarios (DEST_CNPJ distintos)", len({r["DEST_CNPJ"] for r in itens}))
    C("QUANTITATIVA", "EANs validos distintos (EAN_CHAVE_VINCULO)", len({r["EAN_CHAVE_VINCULO"] for r in itens if r["EAN_CHAVE_VINCULO"]}))
    C("QUANTITATIVA", "NCMs distintos", len({r["NCM"] for r in itens}))
    C("QUANTITATIVA", "CFOPs distintos", len({r["CFOP"] for r in itens}), det=" ; ".join("%s (%d)" % kv for kv in Counter(r["CFOP"] for r in itens).most_common()))
    C("QUANTITATIVA", "CST ICMS distintos", len({r["ICMS_CST"] for r in itens if r["ICMS_CST"]}), det=" ; ".join("%s (%d)" % kv for kv in Counter(r["ICMS_CST"] for r in itens if r["ICMS_CST"]).most_common()))
    C("QUANTITATIVA", "CSOSN distintos", len({r["ICMS_CSOSN"] for r in itens if r["ICMS_CSOSN"]}), det=" ; ".join("%s (%d)" % kv for kv in Counter(r["ICMS_CSOSN"] for r in itens if r["ICMS_CSOSN"]).most_common()))
    C("QUANTITATIVA", "Grupos ICMS distintos", len({r["ICMS_GRUPO_TAG"] for r in itens}), det=" ; ".join("%s (%d)" % kv for kv in Counter(vstr(r["ICMS_GRUPO_TAG"]) for r in itens).most_common()))
    C("QUANTITATIVA", "Grupos PIS distintos", len({r["PIS_GRUPO_TAG"] for r in itens}), det=" ; ".join("%s (%d)" % kv for kv in Counter(vstr(r["PIS_GRUPO_TAG"]) + "/" + vstr(r["PIS_CST"]) for r in itens).most_common()))
    C("QUANTITATIVA", "Grupos COFINS distintos", len({r["COFINS_GRUPO_TAG"] for r in itens}), det=" ; ".join("%s (%d)" % kv for kv in Counter(vstr(r["COFINS_GRUPO_TAG"]) + "/" + vstr(r["COFINS_CST"]) for r in itens).most_common()))
    C("QUANTITATIVA", "Grupos IPI distintos", len({r["IPI_GRUPO_TAG"] for r in itens}), det=" ; ".join("%s (%d)" % kv for kv in Counter(vstr(r["IPI_GRUPO_TAG"]) + "/" + vstr(r["IPI_CST"]) for r in itens).most_common()))
    C("QUANTITATIVA", "Itens com grupo IBSCBS", sum(1 for r in itens if r["IBSCBS_PRESENTE"] == "S"))
    C("QUANTITATIVA", "Cenarios fiscais distintos", len(mapa))
    C("QUANTITATIVA", "EANs com divergencia entre documentos", len(eans_div))
    C("QUANTITATIVA", "Registros na amostra", len(amostra))

    # Estruturais
    ids = Counter(r["ID_REGISTRO"] for r in itens)
    dup_ids = [k for k, v in ids.items() if v > 1]
    C("ESTRUTURAL", "ID_REGISTRO duplicado na base", len(dup_ids), "OK" if not dup_ids else "ERRO", ", ".join(dup_ids[:20]))
    sem_vinc = [r for r in itens if not r["ARQ_CAMINHO_COMPLETO"] or not r["CHAVE_NFE"]]
    C("ESTRUTURAL", "Itens sem vinculo ao XML (caminho/chave vazios)", len(sem_vinc), "OK" if not sem_vinc else "ERRO")
    sem_num = [r for r in itens if not r["NUMERO_ITEM"]]
    C("ESTRUTURAL", "Itens sem numero", len(sem_num), "OK" if not sem_num else "ERRO")
    div_qtd = [r for r in canonicos if r["QTD_ITENS_EXTRAIDOS"] != r["QTD_ITENS"]]
    C("ESTRUTURAL", "Notas com itens extraidos <> det no XML", len(div_qtd), "OK" if not div_qtd else "ERRO")
    C("ESTRUTURAL", "Soma QTD_ITENS (det) das notas canonicas = itens extraidos",
      "%d = %d" % (sum(r["QTD_ITENS"] for r in canonicos), len(itens)),
      "OK" if sum(r["QTD_ITENS"] for r in canonicos) == len(itens) else "ERRO")
    # zeros a esquerda: EAN/CNPJ/NCM continuam str e com tamanho esperado
    tipos_ok = all(isinstance(r[c], str) or r[c] is None for r in itens for c in
                   ("EAN", "EAN_TRIB", "EMIT_CNPJ", "DEST_CNPJ", "CHAVE_NFE", "NCM", "CEST", "CFOP", "ICMS_CST", "PIS_CST", "COFINS_CST", "COD_PRODUTO_FORNECEDOR"))
    C("ESTRUTURAL", "Identificadores mantidos como texto (EAN, CNPJ, chave, NCM, CEST, CFOP, CST, cProd)", "SIM" if tipos_ok else "NAO", "OK" if tipos_ok else "ERRO")
    cnpj_bad = [r["ID_REGISTRO"] for r in itens if r["EMIT_CNPJ"] and len(r["EMIT_CNPJ"]) != 14]
    C("ESTRUTURAL", "EMIT_CNPJ com tamanho diferente de 14 (perda de zero?)", len(cnpj_bad), "OK" if not cnpj_bad else "ATENCAO", ", ".join(cnpj_bad[:20]))
    zeros_cnpj = sum(1 for r in itens if (r["EMIT_CNPJ"] or "").startswith("0") or (r["DEST_CNPJ"] or "").startswith("0"))
    C("ESTRUTURAL", "Itens com CNPJ iniciando em zero preservado", zeros_cnpj, "INFO", "Confirma preservacao de zeros a esquerda")
    chave_bad = [r["CHAVE_NFE"] for r in canonicos if r["CHAVE_DV_OK"] != "S"]
    C("ESTRUTURAL", "Chaves com DV invalido", len(chave_bad), "OK" if not chave_bad else "ATENCAO", ", ".join(chave_bad[:10]))
    ean_inv = sum(1 for r in itens if r["EAN_GTIN_VALIDO"] == "N")
    C("ESTRUTURAL", "Itens com EAN ausente/SEM GTIN/invalido", ean_inv, "INFO" if not ean_inv else "ATENCAO")
    criticos = {c: sum(1 for r in itens if not r[c]) for c in ("CFOP", "NCM", "ICMS_GRUPO_TAG", "PIS_GRUPO_TAG", "COFINS_GRUPO_TAG", "DESCRICAO_PRODUTO")}
    for c, n in criticos.items():
        C("ESTRUTURAL", "Itens com %s vazio/ausente" % c, n, "OK" if n == 0 else "ATENCAO")
    C("ESTRUTURAL", "Erros de leitura/parsing", stat.get("ERRO_LEITURA", 0), "OK" if not stat.get("ERRO_LEITURA") else "ATENCAO")
    C("ESTRUTURAL", "Duplicidades (hash identico)", stat.get("DUPLICADO_HASH", 0), "INFO")
    C("ESTRUTURAL", "Duplicidades (mesma chave, conteudo diferente)", stat.get("DUPLICADO_CHAVE", 0), "INFO")
    C("ESTRUTURAL", "Notas com evento de cancelamento encontrado", len({r["CHAVE_NFE"] for r in itens if r["EVENTO_CANCELAMENTO_ENCONTRADO"] == "S"}), "INFO")
    C("ESTRUTURAL", "Tags folha preservadas em 02B_TAGS_ITEM_LONGA", tot_long, "INFO", "Garantia de nao perda de campos nao mapeados")

    # Amostragem
    cen_sem = [m["CENARIO_FISCAL_ID"] for m in mapa if m["QTD_AMOSTRAS_SELECIONADAS"] == 0]
    C("AMOSTRAGEM", "Cenarios sem exemplo na amostra", len(cen_sem), "OK" if not cen_sem else "ERRO", ", ".join(cen_sem))
    cen_1 = [m["CENARIO_FISCAL_ID"] for m in mapa if m["QTD_AMOSTRAS_SELECIONADAS"] == 1 and m["QTD_ITENS"] > 1]
    C("AMOSTRAGEM", "Cenarios com 1 exemplo apesar de >1 item (sem diversidade disponivel de fornecedor/CFOP/NCM/EAN)", len(cen_1), "INFO", ", ".join(cen_1))
    for campo in ("CFOP", "ICMS_GRUPO_TAG", "PIS_GRUPO_TAG", "COFINS_GRUPO_TAG", "IPI_GRUPO_TAG", "EMIT_CNPJ", "ICMS_ORIGEM"):
        base_v = {vstr(r[campo]) for r in itens}
        am_v = {vstr(r[campo]) for r in amostra}
        falt = base_v - am_v
        C("AMOSTRAGEM", "Valores de %s sem exemplo na amostra" % campo, len(falt), "OK" if not falt else "ERRO", ", ".join(sorted(falt)))
    falt_div = set(eans_div) - {r["EAN_CHAVE_VINCULO"] for r in amostra}
    C("AMOSTRAGEM", "EANs divergentes sem exemplo na amostra (teto --max-eans-divergentes-amostra=%d; lista completa em 04)" % a.max_eans_divergentes_amostra,
      len(falt_div), "OK" if not falt_div else ("INFO" if len(eans_div) > a.max_eans_divergentes_amostra else "ERRO"))

    progresso("Passo 4/5 concluidos: divergencias, cenarios, amostra e validacoes")
    # --- Passo 6: escrita
    tipos = {k: v["tipo"] for k, v in COLS.items()}
    inv_cab = [c for c, _ in INV_COLS]
    am_cab = list(AMOSTRA_EXTRA.keys()) + ITEM_COLS
    div_cab = list(divergencias[0].keys()) if divergencias else ["EAN_CHAVE_VINCULO", "DESCRICAO_EXEMPLO", "CAMPO", "QTD_VALORES_DISTINTOS", "VALORES_E_QTD", "ID_REGISTROS_POR_VALOR", "QTD_REGISTROS_EAN", "QTD_FORNECEDORES", "FORNECEDORES", "TIPOS_OPERACAO_ENVOLVIDOS", "DIVERGE_DENTRO_DO_MESMO_FORNECEDOR", "DIVERGE_SO_ENTRE_COMPRAS_DE_TERCEIROS", "STATUS"]
    map_cab = list(mapa[0].keys()) if mapa else ["CENARIO_FISCAL_ID"]
    obs_cab = ["GRUPO_CAMINHO", "TAG", "OCORRENCIAS", "EXEMPLO_ID_REGISTRO"]
    chk_cab = ["CATEGORIA", "CHECAGEM", "RESULTADO", "STATUS", "DETALHE"]
    err_cab = ["NIVEL", "ETAPA", "ARQUIVO", "ERRO", "CAUSA_PROVAVEL", "IMPACTO", "ACAO_TOMADA"]

    # Dicionario (com exemplo real)
    def exemplo(c):
        for r in itens:
            if r.get(c) not in (None, ""):
                return vstr(r[c])[:120]
        return ""
    dic = []
    for c, m in COLS.items():
        nat = m["natureza"]
        dic.append(OrderedDict([
            ("TABELA", "02_BASE_ITENS_XML / 03_AMOSTRA_FISCAL"), ("COLUNA", c), ("DESCRICAO", m["desc"]),
            ("ORIGEM", {"ORIGINAL": "XML", "CONTROLE": "CONTROLE/RASTREABILIDADE", "CALCULADO": "CALCULADO",
                        "CLASSIF_DOC": "CLASSIFICACAO DOCUMENTAL", "INDICIO": "HIPOTESE/INDICIO"}[nat]),
            ("TAG_XML_ORIGEM", m["tag"]), ("TIPO_DADO", m["tipo"]), ("EXEMPLO", exemplo(c)),
            ("DADO_ORIGINAL", fmt_flag(nat == "ORIGINAL")), ("CAMPO_CALCULADO", fmt_flag(nat == "CALCULADO")),
            ("CLASSIFICACAO_DOCUMENTAL", fmt_flag(nat == "CLASSIF_DOC")),
            ("DEPENDE_INTERPRETACAO_FISCAL", fmt_flag(nat == "INDICIO")), ("FORMULA", m["formula"]),
        ]))
    for c, d in AMOSTRA_EXTRA.items():
        dic.append(OrderedDict([("TABELA", "03_AMOSTRA_FISCAL"), ("COLUNA", c), ("DESCRICAO", d), ("ORIGEM", "CONTROLE"),
                                ("TAG_XML_ORIGEM", "-"), ("TIPO_DADO", "TEXTO"), ("EXEMPLO", ""), ("DADO_ORIGINAL", "N"),
                                ("CAMPO_CALCULADO", "N"), ("CLASSIFICACAO_DOCUMENTAL", "N"), ("DEPENDE_INTERPRETACAO_FISCAL", "N"), ("FORMULA", "")]))
    for c, d in INV_COLS:
        dic.append(OrderedDict([("TABELA", "01_INVENTARIO_XML"), ("COLUNA", c), ("DESCRICAO", d), ("ORIGEM", "XML/SISTEMA DE ARQUIVOS"),
                                ("TAG_XML_ORIGEM", "-"), ("TIPO_DADO", "NUMERO" if c in ("TAMANHO_BYTES", "QTD_ITENS", "QTD_ITENS_EXTRAIDOS") else "TEXTO"),
                                ("EXEMPLO", ""), ("DADO_ORIGINAL", "S" if c in ("CHAVE_NFE", "MODELO", "NUMERO_NF", "SERIE", "DATA_EMISSAO", "EMIT_CNPJ", "EMIT_RAZAO_SOCIAL", "DEST_CNPJ", "DEST_RAZAO_SOCIAL", "NATUREZA_OPERACAO", "PROT_CSTAT", "EVENTO_TIPO") else "N"),
                                ("CAMPO_CALCULADO", "N"), ("CLASSIFICACAO_DOCUMENTAL", "N"), ("DEPENDE_INTERPRETACAO_FISCAL", "N"), ("FORMULA", "")]))
    dic_cab = list(dic[0].keys())
    for d in dic:
        tipos[d["COLUNA"]] = tipos.get(d["COLUNA"], d["TIPO_DADO"])
    for c in inv_cab:
        if c in ("TAMANHO_BYTES", "QTD_ITENS", "QTD_ITENS_EXTRAIDOS"):
            tipos[c] = "NUMERO"
        else:
            tipos.setdefault(c, "TEXTO")
            if c in ("CHAVE_NFE", "EMIT_CNPJ", "DEST_CNPJ", "NUMERO_NF", "SERIE", "MODELO", "PROT_CSTAT", "EVENTO_TIPO"):
                tipos[c] = "TEXTO"

    # CSVs
    escrever_csv(os.path.join(exec_dir, "01_INVENTARIO_XML.csv"), inv_cab, inv)
    escrever_csv(os.path.join(exec_dir, "02_BASE_ITENS_XML.csv"), ITEM_COLS, itens)
    escrever_csv(os.path.join(exec_dir, "03_AMOSTRA_FISCAL.csv"), am_cab, amostra)
    escrever_csv(os.path.join(exec_dir, "04_DIVERGENCIAS_PRODUTOS.csv"), div_cab, divergencias)
    escrever_csv(os.path.join(exec_dir, "05_MAPA_CENARIOS_FISCAIS.csv"), map_cab, mapa)
    escrever_csv(os.path.join(exec_dir, "05B_GRUPOS_TAGS_OBSERVADOS.csv"), obs_cab, grupos_obs)
    escrever_csv(os.path.join(exec_dir, "06_LOG_VALIDACAO.csv"), chk_cab, chk)
    escrever_csv(os.path.join(exec_dir, "06B_LOG_ERROS.csv"), err_cab, log.erros)
    escrever_csv(os.path.join(exec_dir, "07_DICIONARIO_CAMPOS.csv"), dic_cab, dic)

    progresso("CSVs gravados")
    if not a.sem_xlsx:
        escrever_xlsx(os.path.join(exec_dir, "BASE_FISCAL_XML.xlsx"), [
            ("01_INVENTARIO_XML", inv_cab, inv),
            ("02_BASE_ITENS_XML", ITEM_COLS, itens),
            ("03_AMOSTRA_FISCAL", am_cab, amostra),
            ("04_DIVERGENCIAS_PRODUTOS", div_cab, divergencias),
            ("05_MAPA_CENARIOS_FISCAIS", map_cab, mapa),
            ("05B_GRUPOS_TAGS_OBSERVADOS", obs_cab, grupos_obs),
            ("06_LOG_VALIDACAO", chk_cab, chk),
            ("06B_LOG_ERROS", err_cab, log.erros),
            ("07_DICIONARIO_CAMPOS", dic_cab, dic),
        ], log, tipos)

    fim = dt.datetime.now()
    execucao = OrderedDict([
        ("ID_EXECUCAO", os.path.basename(exec_dir)), ("INICIO", data_proc), ("FIM", fim.strftime("%Y-%m-%d %H:%M:%S")),
        ("ETAPA", "BASE_FISCAL_AMOSTRAS"), ("VERSAO_EXTRATOR", VERSAO_EXTRATOR), ("ORIGEM", origem),
        ("ARQUIVOS_LIDOS", len(inv)), ("PROCESSADOS", len(canonicos)),
        ("DUPLICADOS", stat.get("DUPLICADO_HASH", 0) + stat.get("DUPLICADO_CHAVE", 0)),
        ("ERROS", stat.get("ERRO_LEITURA", 0)), ("ITENS", len(itens)), ("CENARIOS", len(mapa)), ("AMOSTRAS", len(amostra)),
        ("PARAMETROS", vars(a)),
    ])
    with open(os.path.join(exec_dir, "EXECUCAO.json"), "w", encoding="utf-8") as fh:
        json.dump(execucao, fh, ensure_ascii=False, indent=2)
    print(json.dumps(execucao, ensure_ascii=False, indent=2))
    print("Saida:", exec_dir)
    return 0


if __name__ == "__main__":
    sys.exit(main())
